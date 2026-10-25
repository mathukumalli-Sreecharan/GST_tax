"""Verify semantic cell preservation and the single shared totals result after save."""
from copy import copy
from decimal import Decimal
from gst_calculator import calculate_rate
from io import BytesIO
import posixpath
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter
from monetary_totals import generated_row
from row_splits import RowMapping


def _is_empty_text_normalization(original, exported):
    """OpenPyXL reads <is><t/></is> as '', but saves it as a blank cell.

    This one-way exception applies only to a genuinely empty source text value.
    Whitespace, numbers, errors, formulas, and nonempty text remain strict; all
    style/comment/hyperlink/metadata comparisons still run independently.
    """
    return (original.data_type == 's' and original.value == ''
            and exported.value is None and exported.data_type in {'n', 's', 'inlineStr'})


def validate_export(source, exported, expected, selected, inserted_col, rate_col,
                    recalculate, included_rows, totals, rate_overrides=None, row_splits=None):
    rate_overrides = rate_overrides or {}
    row_splits = row_splits or {}
    row_mapping = RowMapping(row_splits)
    if source.sheetnames != exported.sheetnames:
        raise ValueError('Preservation verification failed: worksheet names changed.')
    total_row = totals['output']['row']
    total_cols = {item['source_column'] + bool(inserted_col and item['source_column'] >= inserted_col)
                  for item in totals['columns'].values() if item['source_column']}
    for old_ws in source:
        new_ws = exported[old_ws.title]
        selected_sheet = old_ws.title == selected
        def mapped(col):
            return col + 1 if selected_sheet and inserted_col and col >= inserted_col else col
        for old in old_ws._cells.values():
            destinations = row_mapping(old.row) if selected_sheet else [old.row]
            for index, destination in enumerate(destinations):
                new = new_ws.cell(destination, mapped(old.column))
                rate_edit = selected_sheet and not inserted_col and (recalculate or old.row in rate_overrides or old.row in row_splits) and old.column == rate_col and old.row in included_rows
                total_edit = selected_sheet and destination == total_row
                split_field = next((key for key, field in totals['columns'].items() if field['source_column'] == old.column), None)
                split_edit = selected_sheet and old.row in row_splits and split_field
                if split_edit:
                    if Decimal(str(new.value)) != Decimal(row_splits[old.row][index][split_field]):
                        raise ValueError('Preservation verification failed: authorized split amount changed.')
                elif not rate_edit and not (total_edit and new.column in total_cols):
                    if old.value != new.value and not _is_empty_text_normalization(old, new):
                        raise ValueError(f'Preservation verification failed: original value at {old_ws.title}!{old.coordinate} changed.')
                if not total_edit and old._style != new._style:
                    raise ValueError(f'Preservation verification failed: style at {old_ws.title}!{old.coordinate} changed.')
                if old.comment != new.comment:
                    raise ValueError('Preservation verification failed: an original cell comment changed.')
                hyperlink = copy(old.hyperlink)
                if hyperlink:
                    hyperlink.ref = new.coordinate
                    # Relationship IDs are regenerated when a link is copied to N rows.
                    # Target, location, display, tooltip and ref remain strict.
                    if row_splits and new.hyperlink:
                        hyperlink.id = new.hyperlink.id
                if hyperlink != new.hyperlink:
                    raise ValueError('Preservation verification failed: an original hyperlink changed.')
        expected_merges = set()
        from openpyxl.utils import get_column_letter
        for area in old_ws.merged_cells.ranges:
            left, top, right, bottom = area.bounds
            if selected_sheet and inserted_col:
                left += left >= inserted_col
                right += right >= inserted_col
            if selected_sheet:
                top = row_mapping(top)[0]
                bottom = row_mapping(bottom)[-1]
            expected_merges.add(f'{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}')
        if expected_merges != {str(area) for area in new_ws.merged_cells.ranges}:
            raise ValueError('Preservation verification failed: original merged regions changed.')
        # Compare serialized geometry/metadata with the explicitly transformed model.
        model = expected[old_ws.title]
        if row_splits:
            for coordinate, cell in new_ws._cells.items():
                if cell.value is not None and coordinate not in model._cells:
                    raise ValueError('Preservation verification failed: unexpected extra output cell or transaction.')
        for attr in ('merged_cells', 'sheet_properties', 'sheet_format', 'sheet_state',
                     'freeze_panes', 'auto_filter', 'print_area', 'print_title_rows', 'print_title_cols',
                     'page_margins', 'page_setup', 'print_options', 'protection', 'sheet_view'):
            if getattr(model, attr) != getattr(new_ws, attr):
                raise ValueError(f'Preservation verification failed: worksheet {attr} changed during export.')
        for key, dimension in model.row_dimensions.items():
            if dict(dimension) != dict(new_ws.row_dimensions[key]) or dimension._style != new_ws.row_dimensions[key]._style:
                raise ValueError('Preservation verification failed: row dimensions changed.')
        for key, dimension in model.column_dimensions.items():
            if dict(dimension) != dict(new_ws.column_dimensions[key]) or dimension._style != new_ws.column_dimensions[key]._style:
                raise ValueError('Preservation verification failed: column dimensions changed.')
    if row_splits and dict(expected.defined_names) != dict(exported.defined_names):
        raise ValueError('Preservation verification failed: transformed named references changed.')
    if source.security != exported.security or source.properties != exported.properties:
        raise ValueError('Preservation verification failed: workbook metadata changed.')
    ws = exported[selected]
    if inserted_col or recalculate or rate_overrides or row_splits:
        old_ws = source[selected]
        for row in included_rows:
            if not inserted_col and not recalculate and row not in rate_overrides and row not in row_splits:
                continue
            calculated = calculate_rate(old_ws.cell(row, totals['columns']['invoice']['source_column']).value,
                                        old_ws.cell(row, totals['columns']['taxable']['source_column']).value)
            if row in row_splits:
                for output_row, part in zip(row_mapping(row), row_splits[row]):
                    if Decimal(str(ws.cell(output_row, rate_col).value)) != Decimal(part['rate']):
                        raise ValueError('Preservation verification failed: authorized split GST rate changed.')
                    for key, field in totals['columns'].items():
                        if field['source_column']:
                            col = field['source_column'] + bool(inserted_col and field['source_column'] >= inserted_col)
                            if Decimal(str(ws.cell(output_row, col).value)) != Decimal(part[key]):
                                raise ValueError('Preservation verification failed: authorized split amount changed.')
                continue
            if ws.cell(row_mapping(row)[0], rate_col).value != rate_overrides.get(row, calculated.rate):
                raise ValueError('Preservation verification failed: exported GST classification differs from the original calculation engine.')
    if row_splits:
        for old_row, parts in row_splits.items():
            for key, field in totals['columns'].items():
                if not field['source_column']:
                    continue
                value = source[selected].cell(old_row, field['source_column']).value
                if value not in (None, '') and sum(Decimal(part[key]) for part in parts) != Decimal(str(value)):
                    raise ValueError('Preservation verification failed: split does not conserve original monetary values.')
    for field in totals['columns'].values():
        if field['source_column'] is None:
            continue
        col = field['source_column']
        col += bool(inserted_col and col >= inserted_col)
        cell = ws.cell(total_row, col)
        if cell.data_type == 'f' or Decimal(str(cell.value)) != Decimal(field['total']):
            raise ValueError('Preservation verification failed: exported currency precision differs from the website total.')
    if ws.cell(total_row, rate_col).value is not None and ws.cell(total_row, rate_col).value != 'GRAND TOTAL':
        raise ValueError('Preservation verification failed: Rate (%) was aggregated.')



def preserve_numeric_cells(source_data, destination, original, selected, inserted_col,
                           rate_col, recalculate, included_rows, totals, rate_overrides=None, row_splits=None):
    """Retain exact XLSX numeric lexemes; OpenPyXL serializes numbers with 16 digits.

    Only existing numeric value elements are restored at explicitly mapped cells.
    Styles, formulas, new rates and intentionally updated generated totals are not
    transplanted. The normal semantic preservation verifier still runs afterwards.
    """
    rate_overrides = rate_overrides or {}
    row_splits = row_splits or {}
    row_mapping = RowMapping(row_splits)
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    rel_ns = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    def parts(archive):
        relationships = {item.get('Id'): item for item in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
        result = {}
        for sheet in ET.fromstring(archive.read('xl/workbook.xml')).find(ns + 'sheets'):
            relationship = relationships[sheet.get(rel_ns + 'id')]
            if relationship.get('TargetMode') == 'External':
                raise ValueError('External worksheet relationships cannot be preserved safely.')
            target = relationship.get('Target', '')
            path = posixpath.normpath(target.lstrip('/') if target.startswith('/') else 'xl/' + target)
            if not path.startswith('xl/') or path not in archive.namelist():
                raise ValueError('Unsupported worksheet package relationship; no output was generated.')
            result[sheet.get('name')] = path
        return result
    replacements = {}
    with ZipFile(BytesIO(source_data)) as source, ZipFile(destination) as exported:
        source_parts, output_parts = parts(source), parts(exported)
        for title, source_part in source_parts.items():
            root = ET.fromstring(exported.read(output_parts[title]))
            cells = {cell.get('r'): cell for cell in root.iter(ns + 'c')}
            monetary_cols = {item['source_column'] for item in totals['columns'].values() if item['source_column']}
            for old in ET.fromstring(source.read(source_part)).iter(ns + 'c'):
                value = old.find(ns + 'v')
                if old.get('t') not in (None, 'n') or value is None or value.text is None or old.find(ns + 'f') is not None:
                    continue
                letter, row = coordinate_from_string(old.get('r'))
                col = column_index_from_string(letter)
                if title == selected:
                    if not inserted_col and (recalculate or row in rate_overrides or row in row_splits) and col == rate_col and row in included_rows:
                        continue
                    if (generated_row(original[title], row) or row in row_splits) and col in monetary_cols:
                        continue
                mapped = col + bool(title == selected and inserted_col and col >= inserted_col)
                for output_row in (row_mapping(row) if title == selected else [row]):
                    coordinate = f'{get_column_letter(mapped)}{output_row}'
                    cell = cells.get(coordinate)
                    if cell is None or cell.get('t') not in (None, 'n') or cell.find(ns + 'f') is not None:
                        raise ValueError(f'Original numeric cell {title}!{old.get("r")} cannot be safely retained at {coordinate}.')
                    new_value = cell.find(ns + 'v')
                    if new_value is None:
                        raise ValueError(f'Original numeric cell {title}!{old.get("r")} lost its value at {coordinate}.')
                    new_value.text = value.text
            replacements[output_parts[title]] = ET.tostring(root, encoding='utf-8')
        entries = [(entry, exported.read(entry.filename)) for entry in exported.infolist()]
    with ZipFile(destination, 'w') as exported:
        for entry, content in entries:
            exported.writestr(entry, replacements.get(entry.filename, content))
