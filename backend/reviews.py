"""Private, expiring review sessions; analysis never writes an Excel output."""
import json
import asyncio
import os
import re
import sqlite3
import time
import uuid
import warnings
from contextlib import contextmanager
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, ConfigDict, Field, model_validator
from openpyxl import load_workbook
from backend.approvals import require_approved_user
from excel_processor import inspect_source, _assert_supported, detect_rate_column, detect_columns, PreservationError
from gst_calculator import calculate_rate, as_decimal
from monetary_totals import aggregate, detect_transaction_columns, monetary_value
from gst_splits import calculate_split
from backend.review_store import durable, SupabaseReviews

router = APIRouter(prefix='/api/excel', tags=['verification'])
TTL = 24 * 3600


@contextmanager
def connect():
    path = Path(os.getenv('REVIEW_STORE_PATH', str(Path(__file__).parent / '.tmp' / 'reviews.sqlite3')))
    path.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(path, timeout=5)
    db.row_factory = sqlite3.Row
    try:
        # Serialize schema inspection/upgrade as well as user mutations. Two first
        # requests must not both inspect an old schema and add the same lease column.
        db.execute('begin immediate')
        db.execute('create table if not exists reviews(id text primary key, owner text not null, filename text not null, source blob not null, options text not null, analysis text not null, decisions text not null, version integer not null, state text not null, result text, created real not null, expires real not null, lease_token text, lease_until real)')
        names={row[1] for row in db.execute('pragma table_info(reviews)')}
        for name,kind in [('lease_token','text'),('lease_until','real')]:
            if name not in names:db.execute(f'alter table reviews add column {name} {kind}')
        db.execute("delete from reviews where expires<=? and (state<>'generating' or coalesce(lease_until,0)<=?)", (time.time(),time.time()))
        db.commit()
        with db:
            yield db
    finally:
        db.close()


def owned(db, review_id, owner):
    item = db.execute('select * from reviews where id=? and owner=? and expires>?', (review_id, owner, time.time())).fetchone()
    if item is None:
        raise HTTPException(404, 'Review session not found or expired.')
    return dict(item)


@contextmanager
def mutation(review_id, owner):
    """One immutable read/version plus atomic commit, for either repository."""
    if durable():
        repository=SupabaseReviews(); item=repository.get(review_id,owner); previous=deepcopy(item)
        yield item
        if item!=previous:repository.save(item,previous)
    else:
        with connect() as db:
            db.execute('begin immediate')
            item=owned(db,review_id,owner);previous=deepcopy(item)
            yield item
            if item!=previous:
                if item.get('state')=='deleted':
                    db.execute('delete from reviews where id=? and owner=? and version=?',(review_id,owner,previous['version']))
                else:
                    db.execute('update reviews set state=?,version=?,decisions=?,result=?,lease_token=?,lease_until=? where id=? and owner=? and version=?',
                        (item['state'],item['version'],item['decisions'],item['result'],item.get('lease_token'),item.get('lease_until'),review_id,owner,previous['version']))


def read_item(review_id,owner):
    if durable():return SupabaseReviews().get(review_id,owner)
    with connect() as db:return owned(db,review_id,owner)


def needs_split(decision):
    return decision.get('choice') == 'WRONG' and len(decision.get('rates', [])) > 1 and not decision.get('split_confirmed')


def snapshot(item):
    analysis = json.loads(item['analysis'])
    decisions = json.loads(item['decisions'])
    required = [row for row in analysis['rows'] if row['status'] != 'MATCHED' or needs_split(decisions.get(str(row['excel_row']), {}))]
    verified = sum(str(row['excel_row']) in decisions and not needs_split(decisions[str(row['excel_row'])]) for row in required)
    blockers=[]
    if analysis['column_totals'].get('issues'):blockers.append('Invalid monetary cells remain. Correct the original values or mappings and start a new review.')
    if analysis['column_totals'].get('boundary_conflicts'):blockers.append('Transaction boundaries remain ambiguous. Clarify the original invoice identity or structural rows before export.')
    return {**analysis, 'generation_blockers':blockers, 'review_id': item['id'], 'filename': item['filename'], 'version': item['version'],
            'state': item['state'], 'lease_until':item.get('lease_until'), 'expires_at': item['expires'], 'decisions': decisions,
            'progress': {'total': len(analysis['rows']), 'matched': len(analysis['rows']) - len(required),
                         'requiring_review': len(required), 'verified': verified, 'remaining': len(required) - verified},
            'result': json.loads(item['result']) if item['result'] else None}


def analyze(source, options):
    inspect_source(source)
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        wb = load_workbook(BytesIO(source), data_only=False)
    if caught:
        wb.close()
        raise PreservationError("Workbook contains unsupported features reported by the Excel reader. Use an authorized simplified copy.")
    try:
        if options['worksheet'] not in wb.sheetnames:
            raise ValueError('Selected worksheet does not exist.')
        ws = wb[options['worksheet']]
        _assert_supported(BytesIO(source), wb, ws)
        first = options['first_data_row']
        if not 2 <= first <= ws.max_row+1:
            raise ValueError('First data row is outside the worksheet.')
        if any(not 1 <= options[key] <= ws.max_column for key in ('invoice_col','taxable_col')):
            raise ValueError('Selected monetary column is outside the worksheet.')
        if options['invoice_col'] == options['taxable_col']:
            raise ValueError('Invoice Value and Taxable Value must use different columns.')
        if max(options['invoice_header_row'] or first-1, options['taxable_header_row'] or first-1) >= first:
            raise ValueError('Header rows must precede the first data row.')
        identities = detect_transaction_columns(ws, first-1)
        if options['invoice_number_col']:
            if not 1 <= options['invoice_number_col'] <= ws.max_column:
                raise ValueError('Invoice number column is outside the worksheet.')
            identities['invoice_number'] = options['invoice_number_col']
        detected_headers=detect_columns(ws)
        if detected_headers:
            earliest=max(detected_headers['invoice_header_row'],detected_headers['taxable_header_row'])+1
            if first>earliest:
                earlier=aggregate(ws,earliest,options['invoice_col'],options['taxable_col'],options['monetary_columns'],identities)
                if any(row<first for row in earlier['included_rows']):
                    raise ValueError('The selected first data row would omit earlier transaction candidates. Review the header and row mapping; no transactions were silently skipped.')
        totals = aggregate(ws, first, options['invoice_col'], options['taxable_col'], options['monetary_columns'], identities)
        columns = {key: field['source_column'] for key, field in totals['columns'].items()}
        existing = detect_rate_column(ws)
        if existing and existing[0] in columns.values():
            raise ValueError('The existing Rate (%) column cannot also be a monetary column.')
        rows = []
        for index in totals['included_rows']:
            def value(col):
                raw = ws.cell(index, col).value if col else None
                return str(raw) if raw is not None else ''
            calc = calculate_rate(ws.cell(index, columns['invoice']).value, ws.cell(index, columns['taxable']).value)
            assigned = calc.rate
            if existing and not options['recalculate_existing']:
                raw = as_decimal(ws.cell(index, existing[0]).value)
                assigned = float(raw) if raw is not None and 0 <= raw <= 100 else None
            zero_gst = (calc.rate is not None and ws.cell(index, columns['invoice']).value == ws.cell(index, columns['taxable']).value
                        and not existing)
            if zero_gst: assigned = None
            issues = [issue for issue in totals['issues'] if issue['row'] == index]
            reasons = [calc.message] if calc.message else []
            ambiguous_boundary = index in totals.get('boundary_conflicts',[])
            if ambiguous_boundary:reasons.append('Transaction boundary is ambiguous; clarify invoice identity or structural labels in the original workbook before export.')
            if zero_gst:reasons.append('Invoice and taxable values are equal. Confirm the actual zero/exempt rate or additional charges; no rate is assumed.')
            if issues:
                reasons.append('Invalid monetary cells: ' + ', '.join(issue['field'].upper() for issue in issues))
            if existing and assigned != calc.rate:
                reasons.append('Existing rate differs from the calculated rate or is missing.')
            amounts = {key: monetary_value(ws.cell(index, col))[0] if col else None for key, col in columns.items()}
            tax_values = [amounts[key] for key in ('igst','cgst','sgst','cess') if columns[key] and ws.cell(index, columns[key]).value not in (None,'')]
            inconsistent = False
            if tax_values and all(value is not None for value in tax_values) and amounts['invoice'] is not None and amounts['taxable'] is not None:
                inconsistent = abs(amounts['invoice'] - amounts['taxable'] - sum(tax_values, Decimal(0))) >= 1
                if inconsistent:
                    reasons.append('Invoice, taxable value and tax components do not reconcile; verify possible multiple GST slabs.')
            status = 'UNRESOLVED' if assigned is None or calc.rate is None or ambiguous_boundary else 'REVIEW_REQUIRED' if calc.flagged or issues or inconsistent or assigned != calc.rate else 'MATCHED'
            if status != 'MATCHED' and not reasons:
                reasons.append('Confirm the actual GST rate or slabs for this transaction.')
            rows.append({'excel_row': index, **{key: value(col) for key,col in identities.items()},
                         **{('invoice_value' if key=='invoice' else 'taxable_value' if key=='taxable' else key): value(col) for key,col in columns.items()},
                         'rate': assigned, 'inferred_rate': calc.rate, 'status': status, 'reason': '; '.join(reasons)})
        return {'worksheet': ws.title, 'rows': rows, 'excluded_rows': totals['excluded_rows'],
                'column_totals': totals, 'transaction_columns': identities}
    finally:
        wb.close()


@router.post('/process')
@router.post('/reviews')
async def start_review(file: UploadFile=File(...), worksheet: str=Form(...), first_data_row: int=Form(...,ge=2),
                       invoice_col: int=Form(...,ge=1), taxable_col: int=Form(...,ge=1),
                       invoice_number_col: int|None=Form(None,ge=1), monetary_columns: str|None=Form(None),
                       invoice_header_row: int|None=Form(None,ge=1), taxable_header_row: int|None=Form(None,ge=1),
                       recalculate_existing: bool=Form(False), confirm_overwrite: bool=Form(False),
                       owner: str=Depends(require_approved_user)):
    from backend.main import read_upload
    if recalculate_existing and not confirm_overwrite:
        raise HTTPException(400,'Confirm overwrite before recalculating an existing Rate(%) column.')
    filename, source = await read_upload(file)
    filename = re.sub(r'[\x00-\x1f\x7f]', '_', (file.filename or filename).replace('\\','/').rsplit('/',1)[-1])[:250] or filename
    try:
        mappings = json.loads(monetary_columns) if monetary_columns else None
        if mappings is not None and not isinstance(mappings,dict):
            raise ValueError('Monetary column mapping must be a JSON object.')
        options = dict(worksheet=worksheet, first_data_row=first_data_row, invoice_col=invoice_col, taxable_col=taxable_col,
                       invoice_number_col=invoice_number_col, monetary_columns=mappings, invoice_header_row=invoice_header_row,
                       taxable_header_row=taxable_header_row, recalculate_existing=recalculate_existing, confirm_overwrite=confirm_overwrite)
        if len(json.dumps(options).encode())>65536:
            raise HTTPException(413,'Review mapping/options exceed the 64 KiB metadata limit.')
        analysis = analyze(source, options)
    except PreservationError as exc:
        raise HTTPException(422,str(exc)) from exc
    except (ValueError,TypeError) as exc:
        raise HTTPException(400,str(exc)) from exc
    if durable():
        return snapshot(SupabaseReviews().create(str(uuid.uuid4()),owner,filename,source,options,analysis))
    with connect() as db:
        db.execute('begin immediate')
        analysis_json=json.dumps(analysis)
        review_bytes=len(analysis_json.encode()) + len(source)
        if db.execute('select count(*) from reviews where owner=?',(owner,)).fetchone()[0] >= 5 or db.execute('select coalesce(sum(length(source)+length(cast(analysis as blob))),0) from reviews').fetchone()[0] + review_bytes > 200*1024*1024:
            raise HTTPException(429,'Review storage is full. Delete an old review or wait for expiry.')
        if review_bytes > 200*1024*1024:
            raise HTTPException(413,'The review is too large for this server. Use a smaller workbook.')
        review_id = str(uuid.uuid4()); now = time.time()
        db.execute('insert into reviews(id,owner,filename,source,options,analysis,decisions,version,state,result,created,expires) values(?,?,?,?,?,?,?,1,\'open\',null,?,?)',
                   (review_id,owner,filename,source,json.dumps(options),analysis_json,'{}',now,now+TTL))
        return snapshot(owned(db,review_id,owner))


@router.get('/reviews')
def list_reviews(owner: str=Depends(require_approved_user)):
    if durable():return {'reviews':SupabaseReviews().list(owner)}
    with connect() as db:
        return {'reviews': [{'review_id': row['id'], 'filename': row['filename'], 'state': row['state'], 'created_at': row['created']}
                            for row in db.execute('select id,filename,state,created from reviews where owner=? order by created desc',(owner,))]}


@router.get('/reviews/{review_id}')
def get_review(review_id: str, owner: str=Depends(require_approved_user)):
    return snapshot(read_item(review_id,owner))


@router.delete('/reviews/{review_id}')
def delete_review(review_id: str, owner: str=Depends(require_approved_user)):
    with mutation(review_id,owner) as item:
        if item['state']=='generating':
            raise HTTPException(409,'Generation is in progress. Recover an expired lease before deleting.')
        item['state']='deleted';item['version']+=1
        return {'deleted':True}



class Decision(BaseModel):
    model_config = ConfigDict(extra='forbid')
    version: int=Field(ge=1)
    choice: str
    slab_count: int|None=Field(default=None,ge=1,le=10)
    rates: list[Decimal]=Field(default_factory=list,max_length=10)

    @model_validator(mode='after')
    def valid(self):
        if self.choice not in {'CORRECT','WRONG'}:
            raise ValueError('Choose CORRECT or WRONG.')
        if self.choice=='CORRECT' and (self.rates or self.slab_count is not None):
            raise ValueError('CORRECT must not change the assigned rate.')
        if self.choice=='WRONG' and (not self.rates or self.slab_count!=len(self.rates) or len(set(self.rates))!=len(self.rates)):
            raise ValueError('Enter the count and every distinct actual slab rate.')
        if any(not rate.is_finite() or rate<0 or rate>100 or rate.as_tuple().exponent < -2 for rate in self.rates):
            raise ValueError('Slab rates must be from 0 to 100 with at most two decimal places.')
        return self


@router.post('/reviews/{review_id}/rows/{excel_row}')
def save_decision(review_id: str, excel_row: int, decision: Decision, owner: str=Depends(require_approved_user)):
    with mutation(review_id,owner) as item:
        if item['state']!='open' or item['version']!=decision.version:
            raise HTTPException(409,'Review changed or generation has started. Reload the review before editing.')
        row = next((row for row in json.loads(item['analysis'])['rows'] if row['excel_row']==excel_row),None)
        if row is None:
            raise HTTPException(404,'Transaction not found.')
        if row['rate'] is None and decision.choice=='CORRECT':
            raise HTTPException(422,'No rate is assigned. Choose WRONG and enter the actual rate or slabs.')
        decisions = json.loads(item['decisions'])
        decisions[str(excel_row)] = {'choice':decision.choice,'rates':[str(rate) for rate in decision.rates], 'saved_at':time.time()}
        item['decisions']=json.dumps(decisions);item['version']+=1
        return snapshot(item)


class SplitRequest(BaseModel):
    model_config = ConfigDict(extra='forbid')
    version: int = Field(ge=1)
    taxable_allocations: list[Decimal] | None = Field(default=None, min_length=2, max_length=10)
    tax_mode: str | None = None
    cess_allocations: list[Decimal] | None = Field(default=None, min_length=2, max_length=10)


class SplitConfirmation(BaseModel):
    model_config = ConfigDict(extra='forbid')
    version: int = Field(ge=1)
    preview_id: str


def checked_split(view, row_number, decision, allocations=None, mode=None, cesses=None):
    row = next((row for row in view['rows'] if row['excel_row'] == row_number), None)
    if row is None or decision.get('choice') != 'WRONG' or len(decision.get('rates', [])) < 2:
        raise HTTPException(422, 'Save WRONG with multiple selected slabs before calculating a split.')
    try:
        if any(issue['row'] == row_number for issue in view['column_totals']['issues']):
            raise ValueError('This transaction has invalid monetary cells. Correct the source monetary values before allocating slabs.')
        preview = calculate_split(row, decision['rates'], allocations, mode, cesses)
        # Nonzero allocations need a mapped destination; never drop an unmapped tax.
        if not preview['manual_required']:
            for key, value in preview['totals'].items():
                if Decimal(value) and not view['column_totals']['columns'][key]['source_column']:
                    raise ValueError('A nonzero ' + key.upper() + ' allocation has no mapped column. Correct the mapping before splitting.')
        return preview
    except (ValueError, InvalidOperation) as exc:
        raise HTTPException(422, str(exc)) from exc


@router.post('/reviews/{review_id}/rows/{excel_row}/split-preview')
def preview_split(review_id: str, excel_row: int, request: SplitRequest, owner: str=Depends(require_approved_user)):
    with mutation(review_id,owner) as item:
        if item['state'] != 'open' or item['version'] != request.version:
            raise HTTPException(409, 'Review changed. Reload before calculating a split.')
        view = snapshot(item); decisions = view['decisions']
        decision = decisions.get(str(excel_row), {})
        preview = checked_split(view, excel_row, decision, request.taxable_allocations, request.tax_mode, request.cess_allocations)
        decision.pop('split_confirmed', None)
        decision.pop('split_confirmed_at', None)
        decision['split_preview'] = {**preview, 'preview_id': str(uuid.uuid4()),
                                    'taxable_allocations': [str(v) for v in request.taxable_allocations] if request.taxable_allocations is not None else None,
                                    'requested_mode': request.tax_mode,
                                    'cess_allocations': [str(v) for v in request.cess_allocations] if request.cess_allocations is not None else None}
        decisions[str(excel_row)] = decision
        item['decisions']=json.dumps(decisions);item['version']+=1
        return snapshot(item)


@router.post('/reviews/{review_id}/rows/{excel_row}/split-confirm')
def confirm_split(review_id: str, excel_row: int, request: SplitConfirmation, owner: str=Depends(require_approved_user)):
    with mutation(review_id,owner) as item:
        if item['state'] != 'open' or item['version'] != request.version:
            raise HTTPException(409, 'Review changed. Reload and confirm the current preview.')
        view = snapshot(item); decisions = view['decisions']
        decision = decisions.get(str(excel_row), {}); preview = decision.get('split_preview', {})
        if preview.get('preview_id') != request.preview_id or preview.get('manual_required', True):
            raise HTTPException(422, 'Calculate a valid replacement-row preview before confirming.')
        checked = checked_split(view, excel_row, decision, preview['taxable_allocations'], preview['requested_mode'], preview.get('cess_allocations'))
        if checked.get('rows') != preview['rows']:
            raise HTTPException(409, 'Split changed. Calculate a fresh preview.')
        decision['split_confirmed'] = True
        decision['split_confirmed_at'] = time.time()
        item['decisions']=json.dumps(decisions);item['version']+=1
        return snapshot(item)


@router.post('/reviews/{review_id}/generate')
async def generate_review(review_id: str, owner: str=Depends(require_approved_user)):
    with mutation(review_id,owner) as item:
        view = snapshot(item)
        if item['state']=='generated':
            return view['result']
        if view['generation_blockers']:
            raise HTTPException(422,' '.join(view['generation_blockers']))
        if item['state']!='open' or view['progress']['remaining']:
            raise HTTPException(409,'Complete all required transaction reviews before generating Excel.')
        item['state']='generating';item['lease_token']=str(uuid.uuid4());item['lease_until']=time.time()+600;item['version']+=1
    from backend.main import generate_excel
    try:
        options = json.loads(item['options'])
        mappings=options.pop('monetary_columns')
        upload=UploadFile(filename=item['filename'],file=BytesIO(SupabaseReviews().source(item) if durable() else item['source']))
        overrides = {}
        splits = {}
        for key, decision in view['decisions'].items():
            if decision['choice']=='WRONG':
                rates=decision['rates']
                if len(rates) == 1:
                    overrides[int(key)] = float(Decimal(rates[0]))
                else:
                    preview = decision.get('split_preview', {})
                    checked = checked_split(view, int(key), decision, preview.get('taxable_allocations'), preview.get('requested_mode'), preview.get('cess_allocations'))
                    if not decision.get('split_confirmed') or checked.get('rows') != preview.get('rows'):
                        raise HTTPException(409, 'Confirm all financial splits before generating Excel.')
                    splits[int(key)] = checked['rows']
                    overrides[int(key)] = float(Decimal(rates[0]))
        result = await generate_excel(file=upload,user_id=owner,monetary_columns=json.dumps(mappings) if mappings is not None else None,
                                      review_context={'transaction_columns':view['transaction_columns'], 'overrides':overrides, 'splits':splits,'defer_history':durable(),
                                                      'durable_review_id':review_id if durable() else None}, **options)
        result['review_id']=review_id
        result['review_complete']=True
        result['summary']['review_complete']=True
        analyzed_rows={row['excel_row']: row for row in view['rows']}
        for row in result['rows']:
            decision=view['decisions'].get(str(row['excel_row']))
            row['review_decision']=decision
            row['inferred_rate']=row['rate']
            row['rate']=(', '.join(part['rate']+'%' for part in splits[row['excel_row']]) if row['excel_row'] in splits
                         else overrides.get(row['excel_row'],analyzed_rows[row['excel_row']]['rate']))
        result['review_progress']=view['progress']
        result['confirmed_splits']={str(key): values for key,values in splits.items()}
        if durable():
            from backend.main import RESULTS
            entry=RESULTS[result['file_id']]
            result['durable_artifact']=SupabaseReviews().publish(item,entry['output'],entry['report'])
            result['review_download_id']=review_id
        with mutation(review_id,owner) as current:
            if current['state']!='generating' or current.get('lease_token')!=item['lease_token'] or current.get('lease_until',0)<=time.time():
                raise HTTPException(409,'Generation lease expired or was replaced; this output was not published.')
            current['state']='generated';current['result']=json.dumps(result);current['version']+=1
        if durable():
            # History is a separate existing feature, populated only AFTER fencing
            # the durable review artifact. Stale workers cannot publish My Files.
            from backend.main import save_generated_file
            history=await asyncio.to_thread(save_generated_file,result['file_id'],owner,item['filename'],entry['output'],result['summary'],len(result['verification']),output_name=result['filename'])
            result['history']=history
            entry['history_saved']=history['saved']
            try:
                with mutation(review_id,owner) as current:
                    current['result']=json.dumps(result);current['version']+=1
            except HTTPException:
                pass  # Durable export is already committed; metadata can be stale.
        return result
    except Exception:
        try:
            with mutation(review_id,owner) as current:
                if current['state']=='generating' and current.get('lease_token')==item['lease_token']:
                    current['state']='open';current['lease_token']=None;current['lease_until']=None;current['version']+=1
        except HTTPException:
            pass
        raise


@router.post('/reviews/{review_id}/recover')
def recover_review(review_id: str, owner: str=Depends(require_approved_user)):
    with mutation(review_id,owner) as item:
        if item['state']!='generating' or (item.get('lease_until') or 0)>time.time():
            raise HTTPException(409,'Only an expired generation lease can be recovered.')
        # Explicitly fence the old worker. Its token-specific private artifacts
        # remain reserved for cleanup; it cannot finalize or save file history.
        item['state']='open';item['lease_token']=None;item['lease_until']=None;item['version']+=1
        return snapshot(item)


@router.get('/reviews/{review_id}/{kind}')
def durable_download(review_id: str,kind: str,owner: str=Depends(require_approved_user)):
    from fastapi import Response
    from urllib.parse import quote
    if kind not in {'download','report'} or not durable():raise HTTPException(404,'Export not available.')
    item=read_item(review_id,owner)
    if item['state']!='generated':raise HTTPException(409,'Complete generation before downloading.')
    result=json.loads(item['result']);data=SupabaseReviews().download(item,kind)
    filename=result['filename'] if kind=='download' else 'verification.csv'
    return Response(data,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if kind=='download' else 'text/csv',
                    headers={'Cache-Control':'private, no-store','X-Content-Type-Options':'nosniff',
                             'Content-Disposition':"attachment; filename*=UTF-8''"+quote(filename,safe='')})
