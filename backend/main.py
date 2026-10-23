from __future__ import annotations

import os
import asyncio
import logging
from contextlib import asynccontextmanager
import re
import tempfile
import time
import uuid
import warnings
from pathlib import Path
from typing import Any

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi import Depends
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from excel_processor import (PreservationError, detect_columns, detect_rate_column,
                             process_workbook, inspect_source, protection_metadata)
from gst_calculator import SUPPORTED_RATES, calculate_rate, as_decimal
from backend.approvals import require_approved_user, router as approval_router
from backend.notifications import deliver_notifications
from backend.files import router as files_router, save_generated_file, file_download, get_file
from backend.security import UploadLimitMiddleware, validate_upload_package, validate_production_config
from verification import report_csv

ROOT = Path(__file__).resolve().parents[1]
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
RESULT_TTL_SECONDS = 60 * 60
RESULTS: dict[str, dict[str, Any]] = {}
# Private process directory; results are intentionally ephemeral and single-worker.
_STORAGE_HANDLE = tempfile.TemporaryDirectory(prefix="gst-ledger-studio-")
STORAGE = Path(_STORAGE_HANDLE.name)

@asynccontextmanager
async def lifespan(app):
    async def cleanup():
        while True:
            await asyncio.sleep(60)
            purge_expired_results()
    task = asyncio.create_task(cleanup())
    async def notify():
        while True:
            try:
                await asyncio.to_thread(deliver_notifications)
            except Exception:
                logging.getLogger(__name__).warning("Notification worker unavailable; queued requests are retained.")
            await asyncio.sleep(30)
    notification_task = asyncio.create_task(notify())
    try:
        yield
    finally:
        task.cancel()
        notification_task.cancel()
        try:
            await notification_task
        except asyncio.CancelledError:
            pass
        try:
            await task
        except asyncio.CancelledError:
            pass
        for entry in list(RESULTS.values()):
            entry["output"].unlink(missing_ok=True)
            entry["report"].unlink(missing_ok=True)
        RESULTS.clear()

validate_production_config()
app = FastAPI(title="GST Ledger Studio API", version="1.0.0", lifespan=lifespan)
app.include_router(approval_router)
app.include_router(files_router)


def configured_cors_origins(value: str | None = None) -> list[str]:
    """Parse allowed browser origins, trimming whitespace and trailing slashes."""
    raw = value if value is not None else os.getenv(
        "CORS_ORIGINS", "http://localhost:3000,http://127.0.0.1:3000"
    )
    return [origin.strip().rstrip("/") for origin in raw.split(",") if origin.strip()]


app.add_middleware(UploadLimitMiddleware)
app.add_middleware(
    CORSMiddleware,
    allow_origins=configured_cors_origins(),
    allow_credentials=False,
    allow_methods=["GET", "POST", "DELETE"],
    allow_headers=["Content-Type", "Authorization", "apikey"],
)


@app.middleware("http")
async def private_response_headers(request, call_next):
    response = await call_next(request)
    response.headers["Cache-Control"] = "private, no-store"
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "no-referrer"
    if os.getenv("APP_ENV") == "production":
        response.headers["Strict-Transport-Security"] = "max-age=31536000"
    return response


def safe_name(name: str) -> str:
    clean = Path(name or "workbook.xlsx").name
    clean = re.sub(r"[^A-Za-z0-9._ -]", "_", clean)[:180]
    return clean if clean.lower().endswith(".xlsx") else f"{Path(clean).stem}.xlsx"


def purge_expired_results() -> None:
    cutoff = time.time() - RESULT_TTL_SECONDS
    for file_id, entry in list(RESULTS.items()):
        if entry["created"] < cutoff:
            entry["output"].unlink(missing_ok=True)
            entry["report"].unlink(missing_ok=True)
            RESULTS.pop(file_id, None)
    for path in STORAGE.glob("inspect_*"):
        try:
            if path.stat().st_mtime < cutoff:
                path.unlink(missing_ok=True)
        except OSError:
            pass


async def read_upload(upload: UploadFile) -> tuple[str, bytes]:
    incoming = Path(upload.filename or "workbook.xlsx").name
    if Path(incoming).suffix.lower() != ".xlsx":
        raise HTTPException(400, "Only .xlsx workbooks are supported.")
    name = safe_name(incoming)
    content = await upload.read(MAX_UPLOAD_BYTES + 1)
    if not content:
        raise HTTPException(400, "The uploaded workbook is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "Workbook exceeds the 20 MB upload limit.")
    validate_upload_package(content)
    return name, content


def write_temp(name: str, data: bytes) -> Path:
    handle, path = tempfile.mkstemp(prefix="inspect_", suffix="_" + safe_name(name), dir=STORAGE)
    os.close(handle)
    result = Path(path)
    result.write_bytes(data)
    return result


def _column_metadata(ws, detected, rate):
    labels: dict[int, str] = {}
    header_row = detected["invoice_header_row"] if detected else 1
    for col in range(1, ws.max_column + 1):
        value = ws.cell(header_row, col).value
        if value is None and detected:
            value = ws.cell(detected["taxable_header_row"], col).value
        labels[col] = str(value).strip() if value is not None and str(value).strip() else f"Column {get_column_letter(col)}"
    invoice_number = None
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40)):
        for cell in row:
            normalized = re.sub(r"[^a-z0-9]+", " ", str(cell.value or "").lower()).strip()
            if "invoice" in normalized and "number" in normalized:
                invoice_number = {"column": cell.column, "header_row": cell.row}
                break
        if invoice_number:
            break
    return [{"index": col, "letter": get_column_letter(col), "header": labels[col],
             "is_invoice_value": bool(detected and col == detected["invoice_col"]),
             "is_taxable_value": bool(detected and col == detected["taxable_col"]),
             "is_rate": bool(rate and col == rate[0]),
             "is_invoice_number": bool(invoice_number and col == invoice_number["column"])}
            for col in range(1, ws.max_column + 1)], invoice_number


def _inspect_sheet(ws):
    detected = detect_columns(ws)
    rate = detect_rate_column(ws)
    columns, invoice_number = _column_metadata(ws, detected, rate)
    preview = []
    for row_index, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 24), values_only=True), 1):
        preview.append({"excel_row": row_index, "values": [cell for cell in row]})
    return {"name": ws.title, "max_row": ws.max_row, "max_column": ws.max_column,
            "detected": detected, "existing_rate": ({"column": rate[0], "header_row": rate[1]} if rate else None),
            "invoice_number": invoice_number, "columns": columns, "preview": preview}


def _safe_error(exc: Exception) -> HTTPException:
    if isinstance(exc, PreservationError):
        return HTTPException(422, str(exc))
    if isinstance(exc, OSError):
        return HTTPException(400, "Temporary workbook storage could not be accessed. Please retry.")
    if isinstance(exc, ValueError):
        return HTTPException(400, str(exc))
    return HTTPException(400, "Workbook could not be processed. Check its format and selected columns.")


@app.get("/api/health")
def health():
    return {"status": "ok"}


@app.post("/api/excel/inspect")
async def inspect_excel(file: UploadFile = File(...), user_id: str = Depends(require_approved_user)):
    purge_expired_results()
    name, content = await read_upload(file)
    temp = write_temp(name, content)
    try:
        _, source_info = inspect_source(content)
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            workbook = load_workbook(temp, read_only=False, data_only=False, keep_links=True)
        try:
            if caught:
                raise HTTPException(422, "OpenPyXL reported workbook content it cannot safely preserve: " + "; ".join(str(w.message) for w in caught))
            sheets = [_inspect_sheet(ws) for ws in workbook.worksheets]
            return {"filename": name, "size_bytes": len(content), "sheets": sheets,
                    "source_info": source_info, "protection": protection_metadata(workbook),
                    "max_upload_bytes": MAX_UPLOAD_BYTES}
        finally:
            workbook.close()
    except HTTPException:
        raise
    except Exception as exc:
        raise _safe_error(exc) from exc
    finally:
        temp.unlink(missing_ok=True)


@app.post("/api/excel/process")
async def process_excel(
    file: UploadFile = File(...),
    worksheet: str = Form(...),
    first_data_row: int = Form(..., ge=2),
    invoice_col: int = Form(..., ge=1),
    taxable_col: int = Form(..., ge=1),
    invoice_number_col: int | None = Form(None, ge=1),
    invoice_header_row: int | None = Form(None, ge=1),
    taxable_header_row: int | None = Form(None, ge=1),
    recalculate_existing: bool = Form(False),
    confirm_overwrite: bool = Form(False),
    user_id: str = Depends(require_approved_user),
):
    purge_expired_results()
    if invoice_col == taxable_col:
        raise HTTPException(400, "Invoice Value and Taxable Value must use different columns.")
    if recalculate_existing and not confirm_overwrite:
        raise HTTPException(400, "Confirm overwrite before recalculating an existing Rate(%) column.")
    name, content = await read_upload(file)
    source = write_temp(name, content)
    output = STORAGE / f"{uuid.uuid4().hex}_{'updated_' + name}"
    original = None
    report_path = None
    try:
        inspect_source(content)
        workbook = load_workbook(source, read_only=False, data_only=False, keep_links=True)
        try:
            if worksheet not in workbook.sheetnames:
                raise HTTPException(400, "Selected worksheet does not exist in this workbook.")
            ws = workbook[worksheet]
            existing = detect_rate_column(ws)
            if existing and recalculate_existing and not confirm_overwrite:
                raise HTTPException(400, "Confirm overwrite before recalculating an existing Rate(%) column.")
            if existing and existing[0] == invoice_col or existing and existing[0] == taxable_col:
                raise HTTPException(400, "The existing Rate(%) column cannot also be an input value column.")
            inv_header = invoice_header_row or max(1, first_data_row - 1)
            tax_header = taxable_header_row or max(1, first_data_row - 1)
            result = process_workbook(source, output, worksheet, invoice_col, taxable_col,
                                      first_data_row - 1, inv_header, tax_header,
                                      existing[0] if existing else None,
                                      bool(existing and recalculate_existing and confirm_overwrite))
            # Read source values and produce a typed row result table from the same calculation engine.
            original = load_workbook(source, read_only=True, data_only=False, keep_links=True)
            original_ws = original[worksheet]
            rate_col = existing[0] if existing else taxable_col + 1
            rows = []
            matches = 0
            for row_index in range(first_data_row, original_ws.max_row + 1):
                invoice_raw = original_ws.cell(row_index, invoice_col).value
                taxable_raw = original_ws.cell(row_index, taxable_col).value
                number = original_ws.cell(row_index, invoice_number_col).value if invoice_number_col else None
                # Ignore fully blank formatted rows in the UI; they remain untouched in the workbook.
                if invoice_raw is None and taxable_raw is None and number is None:
                    continue
                calc = calculate_rate(invoice_raw, taxable_raw)
                inv = as_decimal(invoice_raw)
                taxable_decimal = as_decimal(taxable_raw)
                expected = {rate: inv * 100 / (100 + rate) if inv is not None else None
                            for rate in SUPPORTED_RATES}
                differences = {rate: abs(taxable_decimal - value)
                               if taxable_decimal is not None and value is not None else None
                               for rate, value in expected.items()}
                status = ("AMBIGUOUS" if "ambiguous" in calc.message.lower() else
                          "INVALID" if calc.rate is None and (inv is None or as_decimal(taxable_raw) is None or inv <= 0 or as_decimal(taxable_raw) <= 0 or inv < as_decimal(taxable_raw)) else
                          "UNMATCHED" if calc.rate is None else "REVIEW REQUIRED" if calc.flagged else "MATCHED")
                rows.append({"excel_row": row_index, "invoice_number": str(number) if number is not None else "",
                             "invoice_value": str(invoice_raw) if invoice_raw is not None else "",
                             "taxable_value": str(taxable_raw) if taxable_raw is not None else "",
                             "rate": calc.rate, "difference": str(calc.difference) if calc.difference is not None else "",
                             **{f"difference_{rate}": str(differences[rate]) if differences[rate] is not None else ""
                                for rate in SUPPORTED_RATES},
                             **{f"expected_{rate}": str(expected[rate]) if expected[rate] is not None else ""
                                for rate in SUPPORTED_RATES},
                             "status": status, "reason": calc.message,
                             "invoice_column": invoice_col, "invoice_letter": get_column_letter(invoice_col),
                             "invoice_header": str(original_ws.cell(inv_header, invoice_col).value or f"Column {get_column_letter(invoice_col)}"),
                             "taxable_column": taxable_col, "taxable_letter": get_column_letter(taxable_col),
                             "taxable_header": str(original_ws.cell(tax_header, taxable_col).value or f"Column {get_column_letter(taxable_col)}"),
                             "rate_column": rate_col, "rate_letter": get_column_letter(rate_col)})
                if status == "MATCHED":
                    matches += 1
            original.close()
        finally:
            workbook.close()
        status_counts = {status: sum(1 for row in rows if row["status"] == status)
                         for status in ("MATCHED", "REVIEW REQUIRED", "UNMATCHED", "INVALID", "AMBIGUOUS")}
        rate_counts = {str(rate): sum(1 for row in rows if row["rate"] == rate)
                       for rate in SUPPORTED_RATES}
        summary = {"total": len(rows), "matched": matches,
                   "rate_counts": rate_counts,
                   "five_percent": rate_counts["5"],
                   "twelve_percent": rate_counts["12"],
                   "eighteen_percent": rate_counts["18"],
                   "twenty_eight_percent": rate_counts["28"],
                   "review_required": status_counts["REVIEW REQUIRED"] + status_counts["INVALID"] + status_counts["AMBIGUOUS"],
                   "unmatched": status_counts["UNMATCHED"], "status_counts": status_counts}
        flagged_rows = [row for row in rows if row["status"] != "MATCHED"]
        report_data = []
        for row in flagged_rows:
            report_data.append({"row": row["excel_row"], "invoice_letter": row["invoice_letter"],
                                "invoice_col": row["invoice_column"], "invoice_header": row["invoice_header"],
                                "taxable_letter": row["taxable_letter"], "taxable_col": row["taxable_column"],
                                "taxable_header": row["taxable_header"], "rate_letter": row["rate_letter"],
                                "rate_col": row["rate_column"], "difference": row["difference"],
                                "rate": row["rate"], "message": row["reason"]})
        report_path = STORAGE / f"{uuid.uuid4().hex}_report.csv"
        report_path.write_text(report_csv(report_data), encoding="utf-8")
        if not output.exists():
            raise HTTPException(500, "Workbook processor did not produce an output file.")
        file_id = str(uuid.uuid4())
        original_name = (file.filename or "workbook.xlsx").replace("\\", "/").rsplit("/", 1)[-1]
        original_name = re.sub(r"[\x00-\x1f\x7f]", "_", original_name)[:250] or name
        history = await asyncio.to_thread(save_generated_file, file_id, user_id, original_name, output, summary, len(flagged_rows), output_name="updated_" + name)
        RESULTS[file_id] = {"output": output, "report": report_path, "filename": "updated_" + name,
                            "created": time.time(), "summary": summary, "rows": rows,
                            "worksheet": worksheet, "source_filename": name, "owner_id": user_id,
                            "history_saved": history["saved"]}
        return {"file_id": file_id, "filename": "updated_" + name, "worksheet": worksheet,
                "summary": summary, "rows": rows, "verification": flagged_rows, "history": history,
                "rate_column": {"letter": get_column_letter(rate_col), "index": rate_col}}
    except HTTPException:
        output.unlink(missing_ok=True)
        if report_path:
            report_path.unlink(missing_ok=True)
        raise
    except Exception as exc:
        output.unlink(missing_ok=True)
        if report_path:
            report_path.unlink(missing_ok=True)
        raise _safe_error(exc) from exc
    finally:
        if original is not None:
            original.close()
        source.unlink(missing_ok=True)


def _get_result(file_id: str, user_id: str):
    entry = RESULTS.get(file_id)
    if not entry or time.time() - entry["created"] > RESULT_TTL_SECONDS:
        if entry:
            entry["output"].unlink(missing_ok=True)
            entry["report"].unlink(missing_ok=True)
            RESULTS.pop(file_id, None)
        raise HTTPException(404, "This download is no longer available. Process the workbook again.")
    if entry.get("owner_id") != user_id:
        raise HTTPException(404, "This download is not available.")
    return entry


@app.get("/api/excel/download/{file_id}")
def download_excel(file_id: str, user_id: str = Depends(require_approved_user)):
    # Saved results remain downloadable across sessions and server restarts. Do
    # not serve a stale local copy after its history record is deleted or expired.
    entry = RESULTS.get(file_id)
    if entry is None or entry.get("history_saved"):
        return file_download(file_id, user_id)
    entry = _get_result(file_id, user_id)
    return FileResponse(entry["output"], filename=entry["filename"],
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/excel/report/{file_id}")
def download_report(file_id: str, user_id: str = Depends(require_approved_user)):
    entry = _get_result(file_id, user_id)
    if entry.get("history_saved"):
        get_file(file_id, user_id)
    return FileResponse(entry["report"], filename=f"verification_{Path(entry['source_filename']).stem}.csv",
                        media_type="text/csv; charset=utf-8")
