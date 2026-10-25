from review_helpers import reviewed_post
"""Generated layout coverage; no external accounting workbooks or fixed coordinates."""
from decimal import Decimal
from io import BytesIO
import random
from unittest.mock import patch

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.cell import coordinate_from_string

from excel_processor import process_workbook, PreservationError
from gst_calculator import calculate_rate
from monetary_totals import aggregate, generated_row


def generated(seed, optional=True):
    rng=random.Random(seed)
    keys=['invoice','taxable']+(['igst','cgst','sgst','cess'] if optional else [])
    rng.shuffle(keys)
    headings={'invoice':'Invoice Amount','taxable':'Taxable Value','igst':'Integrated Tax',
              'cgst':'Central Tax','sgst':'State / UT Tax','cess':'CESS'}
    wb=Workbook();ws=wb.active;ws.title=f'Client {seed}'
    header=rng.randint(3,9)
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(keys)+4)
    ws.cell(1,1,'Original merged heading')
    ws.cell(header,1,'Invoice Number')
    cols={key:index for index,key in enumerate(keys,2)}
    tail=max(cols.values())+rng.randint(2,5)
    ws.cell(header,tail,'Trailing source data')
    for key,col in cols.items():ws.cell(header,col,headings[key])
    for offset in range(rng.randint(3,7)):
        row=header+1+offset
        sign=-1 if offset==1 else 1
        ws.cell(row,1,f'INV-{seed}-{offset}')
        for key,col in cols.items():
            value=sign*(118 if key=='invoice' else 100 if key=='taxable' else 9 if key in {'cgst','sgst'} else 18 if key=='igst' else .15)
            ws.cell(row,col,value)
            ws.cell(row,col).number_format='0.000'
            ws.cell(row,col).font=Font(name='Arial',bold=True,color='123456')
            ws.cell(row,col).border=Border(bottom=Side(style='thin'))
            ws.cell(row,col).fill=PatternFill('solid',fgColor='F4F4FF')
        ws.cell(row,tail,f'KEEP-{seed}-{offset}')
        ws.row_dimensions[row].height=21+offset
    ws.column_dimensions[get_column_letter(tail)].width=37
    ws.freeze_panes=f'{get_column_letter(tail)}{header+1}'
    ws.auto_filter.ref=f'A{header}:{get_column_letter(tail)}{ws.max_row}'
    ws.print_area=ws.auto_filter.ref
    ws.print_title_rows=f'1:{header}'
    ws.print_title_cols=f'A:{get_column_letter(tail)}'
    ws.cell(header+1,tail).hyperlink='https://example.com/source'
    wb.properties.title='Original metadata'
    wb.create_sheet('Untouched')['C4']='Do not change this sheet'
    return wb,ws,header,cols


def snapshot(wb):
    stream=BytesIO();wb.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize('seed',range(12))
def test_dynamic_layout_values_refs_and_grand_total_match_shared_result(tmp_path,seed):
    wb,ws,header,cols=generated(seed,optional=seed%3!=0)
    data=snapshot(wb);before=load_workbook(BytesIO(data));old=before[ws.title]
    expected=aggregate(old,header+1,cols['invoice'],cols['taxable'])
    last=max(expected['included_rows']);inserted=cols['taxable']+1
    output=tmp_path/'out.xlsx'
    result=process_workbook(data,output,ws.title,cols['invoice'],cols['taxable'],header)
    after=load_workbook(output);actual=after[ws.title]
    assert after.sheetnames==before.sheetnames
    assert result['column_totals']['output']['row']==last+1
    assert actual.max_row==last+1
    for row in old:
        for cell in row:
            new=actual.cell(cell.row,cell.column+(cell.column>=inserted))
            assert new.value==cell.value
            assert new._style==cell._style
    for key,col in cols.items():
        exported=actual.cell(last+1,col+(col>=inserted))
        assert Decimal(str(exported.value))==Decimal(expected['columns'][key]['total'])
        assert exported.font.bold and exported.border.top.style=='thin'
        assert exported.number_format.endswith('0.00')
    assert actual.cell(last+1,inserted).value is None  # Label uses invoice-number column.
    for row in expected['included_rows']:
        assert actual.cell(row,inserted).value==calculate_rate(old.cell(row,cols['invoice']).value,old.cell(row,cols['taxable']).value).rate
    assert list(after['Untouched'].values)==list(before['Untouched'].values)
    freeze_letter, freeze_row = coordinate_from_string(old.freeze_panes)
    tail_col = column_index_from_string(freeze_letter)
    shifted_tail = tail_col + (tail_col >= inserted)
    assert actual.freeze_panes==f'{get_column_letter(shifted_tail)}{freeze_row}'
    assert actual.cell(header+1,shifted_tail).hyperlink.ref==actual.cell(header+1,shifted_tail).coordinate
    again=tmp_path/'again.xlsx'
    process_workbook(output,again,ws.title,cols['invoice']+(cols['invoice']>=inserted),cols['taxable'],header)
    replay=load_workbook(again)
    assert sum(generated_row(replay[ws.title],row) for row in range(1,replay[ws.title].max_row+1))==1
    assert list(replay[ws.title].values)==list(actual.values)
    before.close();after.close();replay.close()


@pytest.mark.parametrize('content',['Note: filing remarks','Subtotal','GRAND TOTAL'])
def test_occupied_total_destination_is_a_conflict_not_data_loss(tmp_path,content):
    wb,ws,header,cols=generated(21)
    ws.cell(ws.max_row+1,1,content)
    data=snapshot(wb);output=tmp_path/'conflict.xlsx'
    with pytest.raises(PreservationError,match='existing content'):
        process_workbook(data,output,ws.title,cols['invoice'],cols['taxable'],header)
    assert not output.exists()


def test_blank_spacer_allows_total_before_notes_and_retains_trailing_formatting(tmp_path):
    wb,ws,header,cols=generated(22)
    last=ws.max_row
    ws.cell(last+3,1,'Notes: retain original text')
    ws.cell(last+9,ws.max_column).fill=PatternFill('solid',fgColor='FFFF00')
    data=snapshot(wb);output=tmp_path/'out.xlsx'
    result=process_workbook(data,output,ws.title,cols['invoice'],cols['taxable'],header)
    after=load_workbook(output)
    assert result['column_totals']['output']['row']==last+1
    assert after[ws.title].cell(last+3,1).value=='Notes: retain original text'
    assert after[ws.title].cell(last+9,ws.max_column+1).fill.fgColor.rgb.endswith('FFFF00')
    after.close()


def test_all_invalid_footer_without_transaction_identity_is_refused(tmp_path):
    wb,ws,header,cols=generated(23)
    ws.cell(ws.max_row+1,cols['invoice'],'Unrelated unlabelled text')
    data=snapshot(wb)
    with pytest.raises(PreservationError,match='ambiguous'):
        process_workbook(data,tmp_path/'bad.xlsx',ws.title,cols['invoice'],cols['taxable'],header)


def test_wrong_existing_rate_selection_cannot_overwrite_source_column(tmp_path):
    wb,ws,header,cols=generated(24)
    with pytest.raises(PreservationError,match='no original column was overwritten'):
        process_workbook(snapshot(wb),tmp_path/'bad.xlsx',ws.title,cols['invoice'],cols['taxable'],header,existing_rate_col=ws.max_column)


def test_export_validation_removes_a_corrupted_export(tmp_path):
    wb,ws,header,cols=generated(25)
    data=snapshot(wb);output=tmp_path/'corrupt.xlsx'
    real_save=Workbook.save
    def corrupt(model,path):
        model[ws.title].cell(header+1,model[ws.title].max_column,'CORRUPTED')
        real_save(model,path)
    with patch.object(Workbook,'save',corrupt):
        with pytest.raises(PreservationError,match='original value'):
            process_workbook(data,output,ws.title,cols['invoice'],cols['taxable'],header)
    assert not output.exists()



def test_existing_generated_row_with_changed_mappings_reports_conflict(tmp_path):
    wb,ws,header,cols=generated(26)
    first=tmp_path/'first.xlsx'
    process_workbook(snapshot(wb),first,ws.title,cols['invoice'],cols['taxable'],header)
    inserted=cols['taxable']+1
    with pytest.raises(PreservationError,match='outside the selected monetary mappings'):
        process_workbook(first,tmp_path/'second.xlsx',ws.title,cols['invoice']+(cols['invoice']>=inserted),
                         cols['taxable'],header,monetary_columns={'cess':None})


def test_api_returns_explicit_verification_conflict_without_download():
    from fastapi.testclient import TestClient
    from backend.main import app
    from backend.approvals import require_approved_user
    wb,ws,header,cols=generated(27)
    ws.cell(ws.max_row+1,1,'Note: occupied destination')
    app.dependency_overrides[require_approved_user]=lambda:'test-user'
    try:
        response=reviewed_post(TestClient(app),'/api/excel/process',files={'file':('generated.xlsx',snapshot(wb))},data={
            'worksheet':ws.title,'first_data_row':header+1,'invoice_col':cols['invoice'],'taxable_col':cols['taxable']})
        assert response.status_code==422
        assert 'contains existing content' in response.json()['detail']
        assert 'file_id' not in response.json()
    finally:app.dependency_overrides.clear()



def test_grouped_dimensions_do_not_get_default_width_overrides(tmp_path):
    wb,ws,header,cols=generated(28)
    start=min(cols.values());end=max(cols.values())
    ws.column_dimensions.group(get_column_letter(start),get_column_letter(end),hidden=True)
    ws.column_dimensions[get_column_letter(start)].width=31
    output=tmp_path/'grouped.xlsx'
    process_workbook(snapshot(wb),output,ws.title,cols['invoice'],cols['taxable'],header)
    after=load_workbook(output)
    inserted=cols['taxable']+1
    dims=after[ws.title].column_dimensions
    for old in range(start,end+1):
        moved=old+(old>=inserted)
        dimension=next(d for d in dims.values() if d.min<=moved<=d.max)
        assert dimension.width==31 and dimension.hidden
    after.close()



def test_transaction_remarks_are_not_treated_as_footer_notes(tmp_path):
    wb,ws,header,cols=generated(29)
    row=header+1
    ws.cell(row,ws.max_column,'Note: this is a transaction adjustment')
    expected=aggregate(ws,header+1,cols['invoice'],cols['taxable'])
    assert row in expected['included_rows']
    result=process_workbook(snapshot(wb),tmp_path/'remarks.xlsx',ws.title,cols['invoice'],cols['taxable'],header)
    assert result['column_totals']['included_rows']==expected['included_rows']



def test_headers_beyond_the_old_scan_window_are_detected_and_reused(tmp_path):
    from excel_processor import detect_columns, detect_rate_column
    wb=Workbook();ws=wb.active;ws.title='Late header'
    header=67
    ws.cell(header,2,'Invoice Value');ws.cell(header,5,'Taxable Value');ws.cell(header,8,'IGST')
    ws.cell(header+1,2,118);ws.cell(header+1,5,100);ws.cell(header+1,8,18)
    detected=detect_columns(ws)
    assert detected['invoice_header_row']==header
    first=tmp_path/'first.xlsx'
    process_workbook(snapshot(wb),first,ws.title,2,5,header)
    loaded=load_workbook(first)
    assert detect_rate_column(loaded[ws.title])==(6,header)
    second=tmp_path/'second.xlsx'
    process_workbook(first,second,ws.title,2,5,header)
    repeated=load_workbook(second)
    assert repeated[ws.title].max_column==loaded[ws.title].max_column
    assert list(repeated[ws.title].values)==list(loaded[ws.title].values)
    loaded.close();repeated.close()
