from review_helpers import reviewed_post
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from excel_processor import process_workbook, detect_columns, PreservationError
from monetary_totals import aggregate, detect_monetary_columns
from backend.main import app
from backend.approvals import require_approved_user


def fixture_workbook(order=None, multirow=False):
    keys = order or ['invoice', 'taxable', 'igst', 'cgst', 'sgst', 'cess']
    labels = {'invoice':'Invoice Value (â‚¹)', 'taxable':'Taxable VALUE', 'igst':'Integrated Tax (â‚¹)',
              'cgst':'Central Tax', 'sgst':'State/UT Tax', 'cess':'Cess'}
    wb = Workbook(); ws = wb.active; ws.title = 'B2B'
    if multirow:
        ws['A1'] = 'GSTR-2B report'; ws.merge_cells('A1:G1')
        ws['A2'] = 'Invoice Details'
    header = 3 if multirow else 1
    ws.cell(header,1,'Invoice Number')
    for col,key in enumerate(keys,2): ws.cell(header,col,labels[key])
    for row, values in enumerate([
        {'invoice':Decimal('118.005'),'taxable':Decimal('100.005'),'igst':18,'cgst':0,'sgst':0,'cess':Decimal('.105')},
        {'invoice':Decimal('-11.8'),'taxable':-10,'igst':Decimal('-1.8'),'cgst':0,'sgst':0,'cess':Decimal('-.01')},
        {'invoice':Decimal('.1'),'taxable':Decimal('.2'),'igst':None,'cgst':Decimal('.1'),'sgst':Decimal('.2'),'cess':None},
    ],header+1):
        ws.cell(row,1,f'INV-{row}')
        for col,key in enumerate(keys,2): ws.cell(row,col,values[key])
    ws.cell(header+1,2).font=Font(bold=True,color='123456')
    ws.cell(header+1,2).fill=PatternFill('solid',fgColor='FFFF00')
    ws.column_dimensions['B'].width=25; ws.row_dimensions[header+1].height=29
    wb.create_sheet('Reference')['A1']='Unchanged'
    return wb, ws, header, keys


def run(wb,ws,header,keys,tmp_path,name='out.xlsx'):
    stream=BytesIO(); wb.save(stream)
    out=tmp_path/name
    result=process_workbook(stream.getvalue(),out,ws.title,keys.index('invoice')+2,keys.index('taxable')+2,header)
    return result,out


@pytest.mark.parametrize('multirow',[False,True])
@pytest.mark.parametrize('order',[
    ['invoice','taxable','igst','cgst','sgst','cess'],
    ['cess','sgst','taxable','cgst','invoice','igst'],
])
def test_all_six_totals_decimals_negatives_multiline_and_preservation(tmp_path,multirow,order):
    wb,ws,header,keys=fixture_workbook(order,multirow)
    before=[list(row) for row in ws.values]
    result,out=run(wb,ws,header,keys,tmp_path)
    totals=result['column_totals']
    assert {key:item['total'] for key,item in totals['columns'].items()} == {
        'invoice':'106.31','taxable':'90.21','igst':'16.20','cgst':'0.10','sgst':'0.20','cess':'0.10'}
    assert totals['columns']['igst']['blank_count']==1
    assert len(totals['included_rows'])==3
    after=load_workbook(out)
    shift=keys.index('taxable')+3
    for r,row in enumerate(before,1):
        for c,value in enumerate(row,1):
            assert after['B2B'].cell(r,c+(c>=shift)).value==(float(value) if isinstance(value,Decimal) else value)
    assert after['B2B'].cell(header+1,2+(2>=shift)).font.bold
    assert after['B2B'].row_dimensions[header+1].height==29
    assert after['Reference']['A1'].value=='Unchanged'
    output=totals['output']
    for key,item in totals['columns'].items():
        if output['mode']=='grand_total_row':
            col=item['source_column']+(item['source_column']>=shift)
            value=after[output['worksheet']].cell(output['row'],col).value
        else:
            row=list(totals['columns']).index(key)+5
            value=after[output['worksheet']].cell(row,2).value
        assert Decimal(str(value))==Decimal(item['total'])
    assert not any(c.data_type=='f' for sheet in after for row in sheet for c in row)
    # Phase 4 applies identical matching to consistent signed credits.
    assert after['B2B'].cell(header+2,result['rate_col']).value==18
    after.close()


@pytest.mark.parametrize('value,fmt',[
    ('123.45','General'),('INV-118','General'),('â‚¹100','General'),(True,'General'),
    (datetime(2026,10,24),'yyyy-mm-dd'),(.18,'0%'),('#VALUE!','General'),
])
def test_invalid_values_never_become_money(value,fmt):
    wb,ws,header,keys=fixture_workbook()
    ws['D2']=value;ws['D2'].number_format=fmt
    result=aggregate(ws,2,2,3)
    assert result['columns']['igst']['invalid_count']==1
    assert result['columns']['igst']['total']=='-1.80'
    assert result['issues'][0]['row']==2
    assert result['issues'][0]['field']=='igst'
    assert all(str(value) not in issue['message'] for issue in result['issues'])


def test_missing_ambiguous_and_manual_columns():
    wb,ws,header,keys=fixture_workbook()
    ws['H1']='IGST'; ws['H2']=999
    detection=detect_monetary_columns(ws,1)
    assert detection['columns']['igst'] is None
    assert 'ambiguous' in ' '.join(detection['warnings'])
    result=aggregate(ws,2,2,3,{'igst':4})
    assert result['columns']['igst']['total']=='16.20'
    ws['F1']='Other';ws['G1']='Other'
    result=aggregate(ws,2,2,3)
    assert result['columns']['sgst']['total'] is None
    assert result['columns']['cess']['total'] is None
    assert any('SGST' in warning for warning in result['warnings'])
    ws['H1']='Invoice Value'
    assert detect_columns(ws) is None
    for overrides in [{'igst':2},{'igst':100},{'igst':True},{'invoice':4}]:
        with pytest.raises(ValueError): aggregate(ws,2,2,3,overrides)


def test_total_subtotal_and_repeated_headers_are_excluded_and_reprocessing_is_stable(tmp_path):
    wb,ws,header,keys=fixture_workbook()
    ws.append([])  # Leave a safe total destination; retain existing totals below it.
    ws.append(['Subtotal',999,999,999,999,999,999])
    ws.append(['Invoice Number','Invoice Value','Taxable Value','IGST','CGST','SGST','Cess'])
    ws.append(['GRAND TOTAL',9999,9999,9999,9999,9999,9999])
    result,out=run(wb,ws,header,keys,tmp_path)
    totals=result['column_totals']
    assert totals['columns']['invoice']['total']=='106.31'
    assert len(totals['excluded_rows'])==3
    assert totals['output']['mode']=='grand_total_row'
    replay=tmp_path/'replay.xlsx'
    again=process_workbook(out,replay,'B2B',2,3,1)
    assert {k:v['total'] for k,v in again['column_totals']['columns'].items()}=={k:v['total'] for k,v in totals['columns'].items()}
    original,repeat=load_workbook(out),load_workbook(replay)
    assert repeat.sheetnames==original.sheetnames
    assert list(repeat['B2B'].values)==list(original['B2B'].values)
    original.close();repeat.close()


def test_generated_grand_total_is_updated_without_double_counting(tmp_path):
    wb,ws,header,keys=fixture_workbook()
    result,out=run(wb,ws,header,keys,tmp_path)
    repeated=tmp_path/'repeated.xlsx'
    replay=process_workbook(out,repeated,'B2B',2,3,1)
    assert replay['column_totals']['columns']['invoice']['total']=='106.31'
    loaded=load_workbook(repeated)
    markers=[cell for row in loaded['B2B'] for cell in row if cell.value=='GRAND TOTAL']
    assert len(markers)==1
    assert loaded['B2B'].max_row==5
    assert loaded['B2B']['A5'].value=='GRAND TOTAL' # No sum of Rate (%).
    loaded.close()


def test_generated_grand_total_uses_explicit_indian_rupee_format(tmp_path):
    wb,ws,header,keys=fixture_workbook()
    _,out=run(wb,ws,header,keys,tmp_path)
    loaded=load_workbook(out)
    total_row=next(row for row in range(1,loaded['B2B'].max_row+1)
                   if loaded['B2B'].cell(row,1).value=='GRAND TOTAL')
    number_format=loaded['B2B'].cell(total_row,2).number_format
    rupee=chr(0x20b9)
    assert number_format == f'"{rupee}" #,##,##0.00;[Red]-"{rupee}" #,##,##0.00'
    assert '\ufffd' not in number_format and 'Ã' not in number_format and 'â' not in number_format
    loaded.close()


def test_protected_structure_fallback_does_not_bypass_security(tmp_path):
    wb,ws,header,keys=fixture_workbook(multirow=True)
    wb.security.lockStructure=True
    result,out=run(wb,ws,header,keys,tmp_path)
    assert result['column_totals']['output']['worksheet']==ws.title
    assert result['column_totals']['output']['mode']=='grand_total_row'
    loaded=load_workbook(out); assert loaded.security.lockStructure
    assert loaded.sheetnames==['B2B','Reference'];loaded.close()


def test_existing_summary_name_is_not_overwritten(tmp_path):
    wb,ws,header,keys=fixture_workbook(multirow=True)
    wb.create_sheet('GST Totals Summary')['A1']='User content'
    result,out=run(wb,ws,header,keys,tmp_path)
    loaded=load_workbook(out)
    assert loaded['GST Totals Summary']['A1'].value=='User content'
    assert result['column_totals']['output']['worksheet']==ws.title
    assert loaded.sheetnames==wb.sheetnames
    loaded.close()


def test_api_totals_excel_and_verification_agree():
    wb,ws,header,keys=fixture_workbook()
    ws['D2']='invalid'
    stream=BytesIO();wb.save(stream)
    app.dependency_overrides[require_approved_user]=lambda:'test-totals-user'
    try:
        client=TestClient(app)
        response=client.post('/api/excel/process',files={'file':('input.xlsx',stream.getvalue())},data={
            'worksheet':'B2B','first_data_row':2,'invoice_col':2,'taxable_col':3})
        assert response.status_code==200,response.text
        body=response.json(); totals=body['column_totals']
        assert len(totals['included_rows'])==3
        assert any('IGST' in row['reason'] for row in body['rows'])
        assert body['generation_blockers']
        assert client.post('/api/excel/reviews/'+body['review_id']+'/generate').status_code==422
        assert 'file_id' not in body
    finally: app.dependency_overrides.clear()


def test_prasuna_fixture_when_available(tmp_path):
    candidates=sorted(Path('.').glob('*.xlsx'))+sorted(Path('tests/fixtures').glob('**/*.xlsx'))
    if not candidates:pytest.skip('No real GST workbook is available locally')
    source=candidates[0]
    wb=load_workbook(source);ws=next(sheet for sheet in wb if detect_columns(sheet))
    detection=detect_columns(ws);header=max(detection['invoice_header_row'],detection['taxable_header_row'])
    expected=aggregate(ws,header+1,detection['invoice_col'],detection['taxable_col'])
    from backend.reviews import analyze
    options=dict(worksheet=ws.title,first_data_row=header+1,invoice_col=detection['invoice_col'],taxable_col=detection['taxable_col'],invoice_number_col=None,monetary_columns=None,invoice_header_row=detection['invoice_header_row'],taxable_header_row=detection['taxable_header_row'],recalculate_existing=False,confirm_overwrite=False)
    result=analyze(source.read_bytes(),options)
    assert result['column_totals']['columns']==expected['columns']
    assert any(row['status']!='MATCHED' for row in result['rows'])
    # Real final export is intentionally pending end-user decisions.
    wb.close()


def test_gstr_2b_reference_positions_and_header_variations(tmp_path):
    wb=Workbook();ws=wb.active;ws.title='GSTR-2B'
    ws.merge_cells('A1:M1');ws['A1']='GSTR-2B'
    ws['F5']='INVOICE\nVALUE';ws['I4']='taxable amount'
    ws['J4']='I.G.S.T.';ws['K4']='CENTRAL TAX';ws['L4']='State-UT Tax';ws['M4']='CESS'
    ws['A6']='Other header'
    for row,sign in [(7,1),(8,-1)]:
        ws.cell(row,1,'Invoice-'+str(row))
        for col,value in [(6,118),(9,100),(10,18),(11,9),(12,9),(13,1)]:ws.cell(row,col,sign*value)
    detection=detect_columns(ws)
    assert detection['invoice_col']==6 and detection['taxable_col']==9
    stream=BytesIO();wb.save(stream)
    result=process_workbook(stream.getvalue(),tmp_path/'gstr.xlsx',ws.title,6,9,6,5,4)
    assert result['column_totals']['included_rows']==[7,8]
    assert all(item['total']=='0.00' for item in result['column_totals']['columns'].values())
    assert result['rate_col']==10
    output=load_workbook(tmp_path/'gstr.xlsx')
    assert output[ws.title]['J7'].value==18
    assert output[ws.title]['K7'].value==18  # Original IGST preserved after column insertion.
    output.close()


@pytest.mark.parametrize('merged',[False,True])
def test_repeated_totals_preserve_sheet_names_containing_apostrophes(tmp_path,merged):
    wb,ws,header,keys=fixture_workbook(multirow=merged)
    ws.title="Client's B2B"
    result,out=run(wb,ws,header,keys,tmp_path)
    replay=tmp_path/'replay.xlsx'
    second=process_workbook(out,replay,ws.title,2,3,header)
    assert second['column_totals']['columns']['invoice']['total']=='106.31'
    a,b=load_workbook(out),load_workbook(replay)
    assert a.sheetnames==b.sheetnames
    assert list(a[ws.title].values)==list(b[ws.title].values)
    a.close();b.close()


@pytest.mark.parametrize('header',['Invoice Number','Invoice No.','IGST (%)','Tax rate'])
def test_identifier_and_percentage_headers_cannot_be_selected_as_money(header):
    wb,ws,_,_=fixture_workbook()
    ws['D1']=header
    with pytest.raises(ValueError,match='cannot be monetary|cannot be a monetary'):
        aggregate(ws,2,2,3,{'igst':4})



def test_vertically_split_headers_and_currency_units_are_not_transaction_rows(tmp_path):
    wb=Workbook();ws=wb.active
    for col,heading in enumerate(['Invoice','Taxable','Integrated','Central','State/UT','Cess'],1):
        ws.cell(1,col,heading)
        ws.cell(2,col,'Value' if col<=2 else 'Tax' if col<=5 else 'Amount')
        ws.cell(3,col,'INR')
    ws.append([118,100,18,0,0,1])
    ws.append([])
    ws.append(['Invoice','Taxable','Integrated','Central','State/UT','Cess'])
    ws.append(['Value','Value','Tax','Tax','Tax','Amount'])
    ws.append(['INR']*6)
    detection=detect_columns(ws)
    assert detection['invoice_header_row']==3
    assert detection['taxable_header_row']==3
    stream=BytesIO();wb.save(stream)
    result=process_workbook(stream.getvalue(),tmp_path/'split.xlsx',ws.title,1,2,3)
    assert result['column_totals']['included_rows']==[4]
    assert result['column_totals']['columns']['igst']['total']=='18.00'
    assert result['column_totals']['columns']['cess']['total']=='1.00'
