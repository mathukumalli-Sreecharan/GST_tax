"""Loopback-only Phase 3 browser/API test server; never a deployment entrypoint.

Production processing functions run unchanged. Test identity and history saver
are replaced in this separate process so no external credentials/services run.
"""
import os
from contextlib import asynccontextmanager
from io import BytesIO
from tempfile import TemporaryDirectory
from zipfile import ZipFile
from pathlib import Path

# Do this BEFORE importing the application. Do not load any .env file.
for key in list(os.environ):
    if key.startswith(('SUPABASE_', 'RESEND_', 'ADMIN_NOTIFICATION_', 'OWNER_')):
        os.environ.pop(key)
os.environ['APP_ENV']='test'

import uvicorn
from fastapi import File, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook
from backend.main import app
import backend.main as main
from backend.approvals import require_approved_user
from backend.reviews import connect
from tests.test_release_audit import synthetic_b2b


@asynccontextmanager
async def local_lifespan(_app):
    yield  # deliberately no notification worker or external maintenance


def serve():
    with TemporaryDirectory(prefix='gst-phase3-browser-') as temp:
        os.environ['REVIEW_STORE_PATH']=os.path.join(temp,'reviews.sqlite3')
        app.router.lifespan_context=local_lifespan
        app.dependency_overrides[require_approved_user]=lambda:'audit-browser-owner'
        main.save_generated_file=lambda *args,**kwargs:{'saved':False,'warning':'Local audit: external history intentionally disabled.'}

        @app.get('/__audit/fixture')
        def fixture():
            return Response(synthetic_b2b()[0],media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        real_sources=sorted(Path(__file__).resolve().parents[1].glob('*.xlsx'))

        @app.get('/__audit/real-fixtures')
        def real_fixtures():
            return {'count':len(real_sources)}

        @app.get('/__audit/real-fixture/{index}')
        def real_fixture(index:int):
            if not 0<=index<len(real_sources):
                from fastapi import HTTPException
                raise HTTPException(404,'Fixture unavailable')
            return Response(real_sources[index].read_bytes(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        @app.post('/__audit/expire/{review_id}')
        def expire(review_id:str):
            with connect() as db:
                db.execute('update reviews set expires=0 where id=? and owner=?',(review_id,'audit-browser-owner'))
            return {'expired':True}

        @app.post('/__audit/validate')
        async def validate(file:UploadFile=File(...)):
            data=await file.read()
            with ZipFile(BytesIO(data)) as archive:valid=archive.testzip() is None
            wb=load_workbook(BytesIO(data));ws=wb['B2B']
            totals={key:str(ws.cell(15,col).value) for key,col in {'invoice':6,'taxable':9,'igst':12,'cgst':13,'sgst':14,'cess':15}.items()}
            result={'zip_valid':valid,'sheets':wb.sheetnames,'totals':totals,
                    'rates':[ws.cell(row,10).value for row in range(7,15)],
                    'q7_mapped_empty':ws['R7'].value is None and ws['R8'].value is None}
            wb.close();return result

        config=uvicorn.Config(app,host='127.0.0.1',port=3198,log_level='error',access_log=False)
        server=uvicorn.Server(config)

        @app.get('/__audit/shutdown')
        def shutdown():
            server.should_exit=True
            return {'stopping':True}

        server.run()


if __name__=='__main__':
    serve()
