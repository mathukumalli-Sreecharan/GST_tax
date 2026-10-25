"""Universal synthetic finance/layout and concurrency acceptance regressions."""
import asyncio
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from io import BytesIO
import random
from threading import Event
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook,load_workbook
from backend.main import app
from backend.approvals import require_approved_user
from backend.reviews import connect
from excel_processor import process_workbook
from gst_calculator import calculate_rate
from gst_splits import calculate_split
from monetary_totals import aggregate
from tests.test_reviews import source,save


@pytest.mark.parametrize('sign',[-1,1])
@pytest.mark.parametrize('slabs',[2,3,4,5])
@pytest.mark.parametrize('intra',[False,True])
def test_varied_layout_signed_n_slab_financial_and_preservation_invariants(tmp_path,sign,slabs,intra):
    seed=slabs*11+intra; rng=random.Random(seed)
    wb=Workbook();ws=wb.active;ws.title='Layout '+str(seed)
    keys=['invoice','taxable','igst','cgst','sgst','cess'];rng.shuffle(keys)
    keys=['id','name']+keys+['extra']
    headers={'id':'Invoice Number','name':'Supplier Name','invoice':'Invoice Value','taxable':'Taxable Value',
             'igst':'IGST','cgst':'CGST','sgst':'SGST/UTGST','cess':'Cess','extra':'Trailing text'}
    header=rng.randint(2,15);row=header+1
    for col,key in enumerate(keys,1):ws.cell(header,col,headers[key])
    rates=[0,5,12,18,28][:slabs]
    taxable=[Decimal(100*(i+1))*sign for i in range(slabs)]
    taxes=[value*rate/100 for value,rate in zip(taxable,rates)]
    cesses=[Decimal(i+1)*sign for i in range(slabs)]
    total=sum(taxable);gst=sum(taxes);cess=sum(cesses)
    values=dict(id='SAME-INVOICE',name='Synthetic supplier',invoice=total+gst+cess,taxable=total,
                igst=0 if intra else gst,cgst=gst/2 if intra else 0,sgst=gst/2 if intra else 0,cess=cess,extra='Unrelated original text')
    for col,key in enumerate(keys,1):ws.cell(row,col,values[key])
    # Repeated invoice numbers are separate logical source rows, not deduplicated.
    for col,key in enumerate(keys,1):ws.cell(row+1,col,values[key] if key in ('id','name','extra') else None)
    wb.create_sheet('Other '+str(seed))['C4']='Keep other sheet'
    stream=BytesIO();wb.save(stream);raw=stream.getvalue();wb.close()
    mapping={key:keys.index(key)+1 for key in ['invoice','taxable','igst','cgst','sgst','cess']}
    before=load_workbook(BytesIO(raw));totals=aggregate(before[ws.title],row,mapping['invoice'],mapping['taxable'])
    assert totals['included_rows']==[row,row+1]
    record={'invoice_value':values['invoice'],'taxable_value':total,**{key:values[key] for key in ['igst','cgst','sgst','cess']}}
    kwargs={'cess_allocations':cesses}
    if slabs>2:kwargs['taxable_allocations']=taxable
    preview=calculate_split(record,rates,**kwargs)
    assert not preview['manual_required'] and len(preview['rows'])==slabs
    assert all(Decimal(part['taxable'])*sign>0 for part in preview['rows'])
    output=tmp_path/'universal.xlsx'
    result=process_workbook(raw,output,ws.title,mapping['invoice'],mapping['taxable'],header,
        transaction_columns={'invoice_number':1,'supplier_name':2},rate_overrides={row:rates[0],row+1:18},row_splits={row:preview['rows']})
    after=load_workbook(output);actual=after[ws.title];inserted=mapping['taxable']+1
    assert result['column_totals']['source_to_output_rows'][str(row)]==list(range(row,row+slabs))
    assert result['column_totals']['output']['row']==row+slabs+1
    for key,col in mapping.items():
        outcol=col+(col>=inserted)
        assert sum(Decimal(str(actual.cell(r,outcol).value or 0)) for r in range(row,row+slabs))==Decimal(str(values[key]))
        assert Decimal(result['column_totals']['columns'][key]['total'])==Decimal(str(values[key]))
    for old in list(before[ws.title]._cells.values()):
        dests=range(row,row+slabs) if old.row==row else [old.row+(slabs-1 if old.row>row else 0)]
        for dest in dests:
            new=actual.cell(dest,old.column+(old.column>=inserted))
            if old.row==row and keys[old.column-1] in mapping:continue
            assert new.value==old.value
            assert new._style==old._style
    assert after['Other '+str(seed)]['C4'].value=='Keep other sheet'
    before.close();after.close()


def test_rounding_residuals_and_very_small_allocations():
    record={'invoice_value':'348','taxable_value':'300','igst':'48','cgst':'0','sgst':'0','cess':'0'}
    first=calculate_split(record,[18,12],['200.001','99.999'])
    second=calculate_split(record,[18,12],['200.001','99.999'])
    assert first==second
    assert [Decimal(p['taxable']) for p in first['rows']]==[200,100]
    tiny={'invoice_value':'.03','taxable_value':'.03','igst':'0','cgst':'0','sgst':'0','cess':'0'}
    preview=calculate_split(tiny,[5,12],['.015','.015'],tax_mode='igst')
    assert sum(Decimal(p['taxable']) for p in preview['rows'])==Decimal('.03')
    assert all(Decimal(p['taxable'])>0 for p in preview['rows'])
    with pytest.raises(ValueError,match='rounds to zero'):
        calculate_split({**tiny,'invoice_value':'.01','taxable_value':'.01'},[5,12],['.005','.005'],tax_mode='igst')
    with pytest.raises(ValueError,match='Exact preservation'):
        calculate_split({**record,'taxable_value':'300.005'},[18,12])


@pytest.mark.parametrize('sign',[-1,1])
def test_signed_matching_keeps_standard_classification_rules(sign):
    result=calculate_rate(Decimal(118)*sign,Decimal(100)*sign)
    assert result.rate==18 and not result.flagged
    with pytest.raises(ValueError,match='Mixed-sign'):
        calculate_split({'invoice_value':-348,'taxable_value':-300,'igst':48},[18,12])
    with pytest.raises(ValueError,match='signs must match'):
        calculate_split({'invoice_value':-348,'taxable_value':-300,'igst':-48},[18,12],[200,-500])


@pytest.mark.parametrize('label',['Notes: private footer','Remarks: end of report','Prepared by accounts','Net total','Sub total','Summary of taxable supplies','Page total'])
def test_structural_rows_in_supplier_columns_never_become_transactions(label):
    wb=Workbook();ws=wb.active
    ws.append(['GSTIN of supplier','Supplier Name','Invoice Number','Invoice Value','Taxable Value'])
    ws.append(['29ABCDE1234F1Z5','Synthetic','INV-1',118,100])
    ws.append(['29ABCDE1234F1Z5','Synthetic','INV-2',None,None])
    ws.append([label,None,None,118 if 'total' in label.lower() else None,100 if 'total' in label.lower() else None])
    result=aggregate(ws,2,4,5)
    assert result['included_rows']==[2,3]
    assert result['columns']['invoice']['total']=='118.00'


def test_generation_claim_blocks_delete_edit_second_export_and_fences_recovery():
    app.dependency_overrides[require_approved_user]=lambda:'concurrent-owner'
    client=TestClient(app);raw,options,header=source(count=1)
    review=client.post('/api/excel/process',files={'file':('synthetic.xlsx',raw)},data=options).json()
    path='/api/excel/reviews/'+review['review_id'];started=Event();release=Event();calls=[]
    async def generator(**kwargs):
        calls.append(kwargs);started.set();await asyncio.to_thread(release.wait,5)
        return {'file_id':'synthetic-result','rows':[{'excel_row':header+1,'rate':18}], 'summary':{},'verification':[]}
    try:
        with patch('backend.main.generate_excel',side_effect=generator),ThreadPoolExecutor(max_workers=1) as pool:
            future=pool.submit(client.post,path+'/generate')
            assert started.wait(3)
            assert client.delete(path).status_code==409
            assert client.post(path+'/generate').status_code==409
            assert save(client,review,header+1,'CORRECT').status_code==409
            assert client.post(path+'/recover').status_code==409
            release.set();response=future.result()
            assert response.status_code==200,response.text
            assert client.post(path+'/generate').json()['file_id']=='synthetic-result'
            assert len(calls)==1
        # An explicitly expired lease can be recovered; its stale token is removed.
        with connect() as db:
            db.execute("update reviews set state='generating',lease_token='old-worker',lease_until=0 where id=?",(review['review_id'],))
        recovered=client.post(path+'/recover')
        assert recovered.status_code==200 and recovered.json()['state']=='open'
        with connect() as db:
            assert db.execute('select lease_token from reviews where id=?',(review['review_id'],)).fetchone()[0] is None
    finally:release.set();app.dependency_overrides.clear()


def test_late_first_row_cannot_silently_omit_original_transactions():
    app.dependency_overrides[require_approved_user]=lambda:'mapping-owner'
    try:
        client=TestClient(app);raw,options,header=source(count=5)
        response=client.post('/api/excel/process',files={'file':('synthetic.xlsx',raw)},data={**options,'first_data_row':header+2})
        assert response.status_code==400
        assert 'omit earlier transaction candidates' in response.json()['detail']
    finally:app.dependency_overrides.clear()


def test_zero_exempt_transaction_requires_user_rate_and_exports_zero_without_counting_five():
    app.dependency_overrides[require_approved_user]=lambda:'zero-owner'
    try:
        wb=Workbook();ws=wb.active
        ws.append(['Invoice Number','Invoice Value','Taxable Value','IGST','CGST','SGST','Cess'])
        ws.append(['ZERO-1',10,10,0,0,0,0]);stream=BytesIO();wb.save(stream);wb.close()
        client=TestClient(app)
        response=client.post('/api/excel/process',files={'file':('synthetic-zero.xlsx',stream.getvalue())},data={'worksheet':ws.title,'first_data_row':2,'invoice_col':2,'taxable_col':3})
        assert response.status_code==200,response.text
        review=response.json();path='/api/excel/reviews/'+review['review_id']
        assert review['rows'][0]['status']=='UNRESOLVED' and review['rows'][0]['rate'] is None
        assert client.post(path+'/generate').status_code==409
        response=save(client,review,2,rates=[0]);assert response.status_code==200
        with patch('backend.main.save_generated_file',return_value={'saved':False}):result=client.post(path+'/generate')
        assert result.status_code==200,result.text
        data=result.json()
        assert data['rows'][0]['rate']==0 and data['summary']['rate_counts']['0']==1
        assert data['summary']['rate_counts']['5']==0
        output=client.get('/api/excel/download/'+data['file_id'])
        exported=load_workbook(BytesIO(output.content));assert exported.active.cell(2,data['rate_column']['index']).value==0
        exported.close()
    finally:app.dependency_overrides.clear()


def test_unmapped_nonempty_invoice_with_blank_money_is_ambiguous_not_discarded():
    wb=Workbook();ws=wb.active
    ws.append(['Unrecognized document identity','Invoice Value','Taxable Value'])
    ws.append(['DOCUMENT-1',None,None])
    result=aggregate(ws,2,2,3)
    assert result['included_rows']==[2] and result['boundary_conflicts']==[2]
    # Explicit user mapping resolves the boundary without changing original data.
    mapped=aggregate(ws,2,2,3,transaction_columns={'invoice_number':1})
    assert mapped['included_rows']==[2] and mapped['boundary_conflicts']==[]
