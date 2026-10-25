from datetime import datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from backend.main import app
from backend.approvals import require_approved_user
from gst_splits import calculate_split
from tests.test_reviews import save


def values(invoice='348', taxable='300', igst='48', cgst='0', sgst='0', cess='0'):
    return dict(invoice_value=invoice, taxable_value=taxable, igst=igst, cgst=cgst, sgst=sgst, cess=cess)


def test_two_slab_solution_and_manual_three_slabs():
    preview = calculate_split(values(), [18, 12])
    assert [Decimal(row['taxable']) for row in preview['rows']] == [200, 100]
    assert sum(Decimal(row['invoice']) for row in preview['rows']) == 348
    assert calculate_split(values(), [5, 12, 18])['manual_required']
    result = calculate_split(values('683', '600', '83'), [5, 12, 18], [100, 200, 300])
    assert len(result['rows']) == 3
    assert result['totals']['invoice'] == '683.00'


def test_intra_state_cess_and_missing_taxable_manual_allocations():
    assert calculate_split(values('351', '300', '0', '24', '24', '3'), [18, 12])['manual_required']
    preview = calculate_split(values('351', '300', '0', '24', '24', '3'), [18, 12], cess_allocations=[2,1])
    assert [Decimal(p['cgst']) for p in preview['rows']] == [18, 6]
    assert [Decimal(p['cess']) for p in preview['rows']] == [2, 1]
    assert calculate_split(values(taxable=''), [18, 12])['manual_required']
    assert calculate_split(values(taxable=''), [18, 12], [200, 100])['totals']['taxable'] == '300.00'


def test_manual_tax_regime_and_rounding_are_explicit_and_conservative():
    row=values(igst='',cgst='',sgst='')
    assert calculate_split(row,[18,12])['manual_required']
    assert calculate_split(row,[18,12],tax_mode='igst')['totals']['igst']=='48.00'
    # A bounded rounding residual may be distributed, but never changes source totals.
    row=values('3.45','3.00','0.45')
    preview=calculate_split(row,[18,12],[Decimal('1.50'),Decimal('1.50')])
    assert preview['totals']['invoice']=='3.45'
    assert sum(Decimal(p['igst']) for p in preview['rows']) == Decimal('.45')
    with pytest.raises(ValueError):calculate_split(values('349'),[18,12])
    with pytest.raises(ValueError):calculate_split(values(cgst='24',sgst='24'),[18,12])
    with pytest.raises(ValueError):calculate_split(values(),[18,12],tax_mode='cgst_sgst')


def test_zero_rate_slab_is_not_omitted():
    preview=calculate_split(values('318','300','18'),[0,18])
    assert len(preview['rows'])==2
    assert [Decimal(p['taxable']) for p in preview['rows']]==[200,100]
    assert preview['rows'][0]['rate']=='0'


@pytest.mark.parametrize('allocations', [[100, 100], [-1, 301], [0, 300], [250, 50]])
def test_invalid_allocations_rejected(allocations):
    with pytest.raises(ValueError):
        calculate_split(values(), [18, 12], allocations)


def fixture(reverse=False, merged_transaction=False):
    wb = Workbook(); ws = wb.active; ws.title = 'Varied layout'
    fields = ['invoice_number', 'invoice_date', 'invoice', 'taxable', 'igst', 'cgst', 'sgst', 'cess', 'supplier', 'extra']
    if reverse:
        fields.reverse()
    labels = {'invoice_number':'Invoice Number','invoice_date':'Invoice Date','invoice':'Invoice Value','taxable':'Taxable Value','igst':'IGST','cgst':'CGST','sgst':'SGST/UTGST','cess':'Cess','supplier':'Supplier Name','extra':'Unmapped text'}
    header = 4 if reverse else 2
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(fields)); ws.cell(1,1,'Original title')
    for col,key in enumerate(fields,1):ws.cell(header,col,labels[key])
    for offset,(inv,tax,gst) in enumerate([(348,300,48),(683,600,83),(118,100,18)],1):
        data = dict(invoice_number=100+offset, invoice_date=datetime(2026,1,offset), invoice=inv, taxable=tax, igst=gst,cgst=0,sgst=0,cess=0,supplier='Shared Supplier', extra='TRAILING-'+str(offset))
        for col,key in enumerate(fields,1):
            cell=ws.cell(header+offset,col,data[key]);cell.fill=PatternFill('solid',fgColor='DCFCE7')
        cell=ws.cell(header+offset, fields.index('extra')+1);cell.hyperlink='https://example.org/invoice'
        ws.row_dimensions[header+offset].height=23
    footer=header+5;ws.merge_cells(start_row=footer,start_column=1,end_row=footer,end_column=2);ws.cell(footer,1,'Notes: original footer')
    if merged_transaction:ws.merge_cells(start_row=header+1,start_column=9,end_row=header+1,end_column=10)
    ws.auto_filter.ref=f'A{header}:J{header+3}';ws.freeze_panes=f'C{header+1}'
    ws.print_area=f'A1:J{footer}'
    wb.create_sheet('Other sheet')['D8']=123.4567890123456
    stream=BytesIO();wb.save(stream);wb.close()
    return stream.getvalue(),dict(worksheet=ws.title,first_data_row=header+1,invoice_col=fields.index('invoice')+1,taxable_col=fields.index('taxable')+1),header,fields


@pytest.fixture
def client():
    app.dependency_overrides[require_approved_user]=lambda:'split-owner'
    yield TestClient(app)
    app.dependency_overrides.clear()


def preview_confirm(client, review, row, allocations=None):
    path=f"/api/excel/reviews/{review['review_id']}/rows/{row}"
    body={'version':review['version']}
    if allocations is not None:body['taxable_allocations']=allocations
    response=client.post(path+'/split-preview',json=body)
    assert response.status_code==200,response.text
    review=response.json();part=review['decisions'][str(row)]['split_preview']
    assert not part['manual_required']
    response=client.post(path+'/split-confirm',json={'version':review['version'],'preview_id':part['preview_id']})
    assert response.status_code==200,response.text
    return response.json()


@pytest.mark.parametrize('reverse',[False,True])
def test_exact_replacement_mapping_totals_metadata_and_reprocessing(client,reverse):
    raw,options,header,fields=fixture(reverse)
    response=client.post('/api/excel/process',files={'file':('layout.xlsx',raw)},data=options)
    assert response.status_code==200,response.text
    review=response.json();path='/api/excel/reviews/'+review['review_id']
    review=save(client,review,header+1,rates=[18,12]).json()
    assert client.post(path+'/generate').status_code==409
    review=preview_confirm(client,review,header+1)
    review=save(client,review,header+2,rates=[5,12,18]).json()
    review=preview_confirm(client,review,header+2,[100,200,300])
    assert review['progress']['remaining']==0
    with patch('backend.main.save_generated_file',return_value={'saved':False}):
        response=client.post(path+'/generate')
    assert response.status_code==200,response.text
    result=response.json();output=client.get('/api/excel/download/'+result['file_id']).content
    wb=load_workbook(BytesIO(output));ws=wb[options['worksheet']]
    mapping=result['column_totals']['source_to_output_rows']
    report=client.get('/api/excel/report/'+result['file_id'])
    assert report.status_code==200
    assert 'Output Excel rows' in report.text and 'Confirmed replacement allocations' in report.text
    assert mapping=={str(header+1):[header+1,header+2],str(header+2):[header+3,header+4,header+5],str(header+3):[header+6]}
    assert result['column_totals']['output']['row']==header+7
    rate=result['rate_column']['index']
    col=lambda key:fields.index(key)+1+int(fields.index(key)+1>=rate)
    assert [ws.cell(r,rate).value for r in range(header+1,header+7)]==[18,12,5,12,18,18]
    assert [ws.cell(r,col('invoice_number')).value for r in range(header+1,header+7)]==[101,101,102,102,102,103]
    assert sum(Decimal(str(ws.cell(r,col('invoice')).value)) for r in range(header+1,header+7))==1149
    for key,total in [('invoice',1149),('taxable',1000),('igst',149),('cgst',0),('sgst',0),('cess',0)]:
        assert ws.cell(header+7,col(key)).value==total
        assert Decimal(result['column_totals']['columns'][key]['total'])==total
    for row in range(header+1,header+7):
        cell=ws.cell(row,col('extra'))
        assert cell.hyperlink.target=='https://example.org/invoice'
        assert cell.fill.fgColor.rgb=='00DCFCE7'
        assert ws.row_dimensions[row].height==23
    assert ws.cell(header+8,1).value=='Notes: original footer'
    assert wb['Other sheet']['D8'].value==123.4567890123456
    assert client.post(path+'/generate').json()['file_id']==result['file_id']
    wb.close()
    # Reprocessing already split output must not add rows or duplicate totals.
    rerun=client.post('/api/excel/process',files={'file':('split.xlsx',output)},data={**options,'invoice_col':col('invoice'),'taxable_col':col('taxable')}).json()
    assert len(rerun['rows'])==6
    for row in rerun['rows']:
        if row['status']!='MATCHED':rerun=save(client,rerun,row['excel_row'],'CORRECT').json()
    with patch('backend.main.save_generated_file',return_value={'saved':False}):
        final=client.post('/api/excel/reviews/'+rerun['review_id']+'/generate')
    assert final.status_code==200,final.text
    assert final.json()['column_totals']['output']['row']==header+7


def test_confirmation_versions_ownership_and_edits_invalidate_split(client):
    raw,options,header,_=fixture()
    review=client.post('/api/excel/process',files={'file':('a.xlsx',raw)},data=options).json()
    review=save(client,review,header+1,rates=[18,12]).json()
    rowpath=f"/api/excel/reviews/{review['review_id']}/rows/{header+1}"
    assert client.post(rowpath+'/split-confirm',json={'version':review['version'],'preview_id':'invented'}).status_code==422
    response=client.post(rowpath+'/split-preview',json={'version':review['version']})
    updated=response.json()
    assert client.post(rowpath+'/split-preview',json={'version':review['version']}).status_code==409
    app.dependency_overrides[require_approved_user]=lambda:'other-owner'
    assert client.post(rowpath+'/split-preview',json={'version':updated['version']}).status_code==404
    app.dependency_overrides[require_approved_user]=lambda:'split-owner'
    confirmed=preview_confirm(client,updated,header+1)
    edited=save(client,confirmed,header+1,rates=[12,18]).json()
    assert not edited['decisions'][str(header+1)].get('split_confirmed')
    assert 'split_preview' not in edited['decisions'][str(header+1)]


def test_unsupported_merged_transaction_fails_without_publishing(client):
    raw,options,header,_=fixture(merged_transaction=True)
    review=client.post('/api/excel/process',files={'file':('a.xlsx',raw)},data=options).json()
    review=save(client,review,header+1,rates=[18,12]).json()
    review=preview_confirm(client,review,header+1)
    for row in review['rows']:
        if row['status']!='MATCHED' and str(row['excel_row']) not in review['decisions']:
            review=save(client,review,row['excel_row'],'CORRECT').json()
    with patch('backend.main.save_generated_file') as publish:
        response=client.post('/api/excel/reviews/'+review['review_id']+'/generate')
    assert response.status_code==422,response.text
    publish.assert_not_called()


@pytest.mark.parametrize('damage', ['shared_text', 'style', 'amount', 'rate', 'later_transaction', 'extra_copy'])
def test_strict_verifier_detects_changes_after_split(tmp_path, damage):
    from excel_processor import process_workbook, PreservationError
    from workbook_preservation import preserve_numeric_cells
    raw, options, header, fields = fixture()
    rate = options['taxable_col'] + 1
    def mapped(key):
        col = fields.index(key) + 1
        return col + (col >= rate)
    output = tmp_path / 'must-not-exist.xlsx'
    def corrupt(*args):
        preserve_numeric_cells(*args)
        wb = load_workbook(output); ws = wb[options['worksheet']]
        if damage == 'shared_text':ws.cell(header+2, mapped('extra'), 'Lost original detail')
        elif damage == 'style':ws.cell(header+2, mapped('extra')).fill = PatternFill('solid', fgColor='FF0000')
        elif damage == 'amount':ws.cell(header+2, mapped('invoice'), 1000)
        elif damage == 'rate':ws.cell(header+2, rate, 28)
        elif damage == 'later_transaction':ws.cell(header+3, mapped('invoice'), 1000)
        else:ws.cell(header+20, mapped('invoice'), 348)
        wb.save(output); wb.close()
    with patch('excel_processor.preserve_numeric_cells', side_effect=corrupt):
        with pytest.raises(PreservationError):
            process_workbook(raw, output, options['worksheet'], options['invoice_col'], options['taxable_col'], header,
                             rate_overrides={header+1:18}, transaction_columns={'invoice_number':1,'invoice_date':2,'supplier_name':9},
                             row_splits={header+1:calculate_split(values(),[18,12])['rows']})
    assert not output.exists()


def test_numeric_lexemes_survive_shared_numeric_ids_and_unrelated_sheet(tmp_path):
    from tests.test_numeric_preservation import raw_values
    from excel_processor import process_workbook
    from zipfile import ZipFile
    raw, options, header, fields = fixture()
    raw = raw_values(raw, {'xl/worksheets/sheet1.xml': {'A'+str(header+1):'12345678901234567'},
                           'xl/worksheets/sheet2.xml': {'D8':'0.30000000000000004'}})
    output = tmp_path / 'exact.xlsx'
    process_workbook(raw, output, options['worksheet'], options['invoice_col'], options['taxable_col'], header,
                     rate_overrides={header+1:18}, transaction_columns={'invoice_number':1,'invoice_date':2,'supplier_name':9},
                     row_splits={header+1:calculate_split(values(),[18,12])['rows']})
    wb=load_workbook(output)
    assert wb[options['worksheet']].cell(header+1,1).value == 12345678901234567
    assert wb[options['worksheet']].cell(header+2,1).value == 12345678901234567
    assert wb['Other sheet']['D8'].value == 0.30000000000000004
    wb.close()
    with ZipFile(output) as archive:
        assert archive.read('xl/worksheets/sheet1.xml').count(b'12345678901234567') == 2


def test_splitting_does_not_recalculate_confirmed_existing_rates(tmp_path):
    from excel_processor import process_workbook
    raw, options, header, fields = fixture()
    wb=load_workbook(BytesIO(raw));ws=wb[options['worksheet']]
    rate=ws.max_column+1;ws.cell(header,rate,'Rate (%)')
    for row in range(header+1,header+4):ws.cell(row,rate,7)
    stream=BytesIO();wb.save(stream);wb.close()
    output=tmp_path/'kept.xlsx'
    process_workbook(stream.getvalue(),output,options['worksheet'],options['invoice_col'],options['taxable_col'],header,
                     existing_rate_col=rate, rate_overrides={header+1:18},
                     transaction_columns={'invoice_number':1,'invoice_date':2,'supplier_name':9},
                     row_splits={header+1:calculate_split(values(),[18,12])['rows']})
    wb=load_workbook(output);ws=wb[options['worksheet']]
    assert [ws.cell(row,rate).value for row in range(header+1,header+5)] == [18,12,7,7]
    wb.close()


def test_manual_cess_confirmation_and_export_preserve_all_components(client):
    raw, options, header, fields=fixture()
    wb=load_workbook(BytesIO(raw));ws=wb[options['worksheet']]
    ws.cell(header+1,fields.index('invoice')+1,351)
    ws.cell(header+1,fields.index('cess')+1,3)
    stream=BytesIO();wb.save(stream);wb.close()
    review=client.post('/api/excel/process',files={'file':('cess.xlsx',stream.getvalue())},data=options).json()
    review=save(client,review,header+1,rates=[18,12]).json()
    path=f"/api/excel/reviews/{review['review_id']}"
    rowpath=f'{path}/rows/{header+1}'
    review=client.post(rowpath+'/split-preview',json={'version':review['version']}).json()
    assert review['decisions'][str(header+1)]['split_preview']['manual_required']
    assert client.post(path+'/generate').status_code==409
    response=client.post(rowpath+'/split-preview',json={'version':review['version'],'cess_allocations':[2,1]})
    assert response.status_code==200,response.text
    review=response.json();preview=review['decisions'][str(header+1)]['split_preview']
    assert [Decimal(part['cess']) for part in preview['rows']]==[2,1]
    review=client.post(rowpath+'/split-confirm',json={'version':review['version'],'preview_id':preview['preview_id']}).json()
    for row in review['rows']:
        if row['status']!='MATCHED' and str(row['excel_row']) not in review['decisions']:
            review=save(client,review,row['excel_row'],'CORRECT').json()
    with patch('backend.main.save_generated_file',return_value={'saved':False}):
        result=client.post(path+'/generate')
    assert result.status_code==200,result.text
    assert Decimal(result.json()['column_totals']['columns']['cess']['total'])==3
