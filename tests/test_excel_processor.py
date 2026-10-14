from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from excel_processor import detect_columns, process_workbook
from excel_processor import detect_rate_column
from gst_calculator import calculate_rate


def test_detection_varied_positions_and_workbook_preservation(tmp_path):
    source, output = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "B2B"
    ws.merge_cells("A1:F1")
    ws["A1"] = "GST report"
    ws["B3"] = "Invoice Value (₹)"
    ws["E3"] = "Taxable Value (₹)"
    ws["B4"] = 105
    ws["E4"] = 100
    ws["B5"] = 118
    ws["E5"] = 100
    ws["B6"] = 112
    ws["E6"] = 100
    ws["B7"] = 128
    ws["E7"] = 100
    ws["A4"] = "kept"
    ws["A4"].font = Font(bold=True)
    ws["A4"].fill = PatternFill("solid", fgColor="FFFF00")
    ws.column_dimensions["A"].width = 27
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = "B4"
    ws.auto_filter.ref = "A3:E5"
    ws.row_dimensions[6].hidden = True
    wb.create_sheet("Other")
    wb.save(source)

    loaded = load_workbook(source)
    found = detect_columns(loaded["B2B"])
    assert found["invoice_col"] == 2 and found["taxable_col"] == 5
    loaded.close()
    result = process_workbook(source, output, "B2B", 2, 5, 3, 3, 3)
    assert result["rate_col"] == 6
    out = load_workbook(output)
    ws = out["B2B"]
    assert out.sheetnames == ["B2B", "Other"]
    assert ws["F3"].value == "Rate(%)"
    assert ws["F4"].value == 5 and ws["F5"].value == 18
    assert ws["F6"].value == 12 and ws["F7"].value == 28
    assert ws["A4"].value == "kept" and ws["A4"].font.bold
    assert ws.column_dimensions["A"].width == 27
    assert ws.row_dimensions[4].height == 30 and ws.row_dimensions[6].hidden
    assert ws.freeze_panes == "B4"
    assert ws.auto_filter.ref == "A3:F5"
    assert "A1:G1" in {str(r) for r in ws.merged_cells.ranges}
    out.close()


def test_does_not_overwrite_input(tmp_path):
    source = tmp_path / "input.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "Invoice Value", "Taxable Value"
    wb.save(source)
    original = source.read_bytes()
    try:
        process_workbook(source, source, "Sheet", 1, 2, 1)
    except ValueError:
        pass
    assert source.read_bytes() == original


def test_existing_rate_column_is_never_duplicated_or_overwritten_without_confirmation(tmp_path):
    source, kept, recalculated = tmp_path / "source.xlsx", tmp_path / "kept.xlsx", tmp_path / "recalculated.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["C1"] = "Invoice Value", "Taxable Value", "Rate(%)"
    ws["A2"], ws["B2"], ws["C2"] = 118, 100, 12
    wb.save(source)
    found = load_workbook(source)
    assert detect_rate_column(found.active) == (3, 1)
    found.close()

    original_bytes = source.read_bytes()
    process_workbook(source, kept, "Sheet", 1, 2, 1, 1, 1)
    assert kept.read_bytes() == original_bytes

    process_workbook(source, recalculated, "Sheet", 1, 2, 1, 1, 1, 3, True)
    out = load_workbook(recalculated)
    assert out.active.max_column == 3
    assert out.active["C2"].value == 18
    out.close()


def test_provided_workbook_regression_pairs(tmp_path):
    pairs = [
        ("KK MAY.xlsx", "updated_KK MAY.xlsx"),
        ("KKENGIN APRIL-26.xlsx", "updated_KKENGIN APRIL-26.xlsx"),
    ]
    for original_name, expected_name in pairs:
        source, expected = Path(original_name), Path(expected_name)
        if not source.exists() or not expected.exists():
            pytest.skip(f"Regression workbooks not available: {original_name}, {expected_name}")
        output = tmp_path / ("processed_" + expected.name)
        original_wb = load_workbook(source, data_only=False)
        original_ws = original_wb["B2B"]
        detection = detect_columns(original_ws)
        result = process_workbook(source, output, "B2B", detection["invoice_col"], detection["taxable_col"],
                                  max(detection["invoice_header_row"], detection["taxable_header_row"]),
                                  detection["invoice_header_row"], detection["taxable_header_row"])
        actual_wb = load_workbook(output, data_only=False)
        expected_wb = load_workbook(expected, data_only=False)
        actual, reference = actual_wb["B2B"], expected_wb["B2B"]

        assert detect_rate_column(actual) == (10, 5)
        # Every original cell/value/style remains at its position, shifted one col after insertion.
        for row in range(1, original_ws.max_row + 1):
            for col in range(1, original_ws.max_column + 1):
                target_col = col if col < 10 else col + 1
                old, new = original_ws.cell(row, col), actual.cell(row, target_col)
                assert (new.value or None) == (old.value or None)
                assert new._style == old._style
            assert actual.row_dimensions[row].height == original_ws.row_dimensions[row].height
            assert actual.row_dimensions[row].hidden == original_ws.row_dimensions[row].hidden
        assert actual.freeze_panes == original_ws.freeze_panes
        assert actual_wb.sheetnames == original_wb.sheetnames
        def width_at(sheet, column):
            return next((dimension.width for dimension in sheet.column_dimensions.values()
                         if dimension.min <= column <= dimension.max), sheet.sheet_format.defaultColWidth)
        def hidden_at(sheet, column):
            return next((dimension.hidden for dimension in sheet.column_dimensions.values()
                         if dimension.min <= column <= dimension.max), False)
        for col in range(1, original_ws.max_column + 1):
            target_col = col if col < 10 else col + 1
            assert width_at(actual, target_col) == width_at(original_ws, col)
            assert hidden_at(actual, target_col) == hidden_at(original_ws, col)
        shifted_merges = set()
        for merged in original_ws.merged_cells.ranges:
            min_col, min_row, max_col, max_row = range_boundaries(str(merged))
            if min_col >= 10:
                min_col += 1
                max_col += 1
            elif max_col >= 10:
                max_col += 1
            shifted_merges.add(f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}")
        assert shifted_merges == {str(r) for r in actual.merged_cells.ranges}
        # Existing processed workbook rate values are the reference classifications.
        expected_flagged_rows = set()
        for row in range(7, original_ws.max_row + 1):
            calculated = calculate_rate(original_ws.cell(row, detection["invoice_col"]).value,
                                        original_ws.cell(row, detection["taxable_col"]).value)
            assert actual.cell(row, 10).value == calculated.rate
            assert actual.cell(row, 10).value == reference.cell(row, 10).value
            if calculated.flagged:
                expected_flagged_rows.add(row)
        # Every reported row has a reason, and unmatched classifications are blank.
        flagged = {item["row"]: item for item in result["flagged"]}
        assert set(flagged) == expected_flagged_rows
        assert all("message" in item and item["rate_letter"] == "J" for item in flagged.values())
        assert all(actual.cell(row, 10).value is None for row in flagged if flagged[row]["rate"] is None)
        # Processing a workbook with Rate(%) again keeps the package untouched by default.
        kept = tmp_path / ("kept_" + expected.name)
        process_workbook(output, kept, "B2B", 6, 9, 6, 6, 5)
        assert kept.read_bytes() == output.read_bytes()
        original_wb.close()
        actual_wb.close()
        expected_wb.close()


def test_formula_workbook_stops_instead_of_risking_bad_references(tmp_path):
    source, output = tmp_path / "formula.xlsx", tmp_path / "output.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "Invoice Value", "Taxable Value"
    ws["C2"] = "=A2+B2"
    wb.save(source)
    from excel_processor import PreservationError
    with pytest.raises(PreservationError, match="contains formulas"):
        process_workbook(source, output, "Sheet", 1, 2, 1, 1, 1)
    assert not output.exists()
