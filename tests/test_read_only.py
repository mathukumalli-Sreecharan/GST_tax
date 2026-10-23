"""Read-only inputs are snapshots, never edited source files."""
from io import BytesIO
import os
import stat
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Border, Side, PatternFill
from openpyxl.workbook.protection import WorkbookProtection

from excel_processor import PreservationError, inspect_source, process_workbook


def workbook_bytes(*, sheet_protected=False, structure=False, existing=False):
    wb = Workbook()
    ws = wb.active
    ws.append(["Invoice Value", "Taxable Value"] + (["Rate(%)"] if existing else []))
    for rate in (5, 12, 18, 28):
        ws.append([100 + rate, 100] + ([rate] if existing else []))
    ws["A2"].font = Font(bold=True, color="112233")
    ws["A2"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["A2"].border = Border(bottom=Side(style="thin"))
    ws["A2"].number_format = "0.00"
    ws.row_dimensions[2].height = 30
    ws.column_dimensions["A"].width = 25
    ws.protection.sheet = sheet_protected
    if sheet_protected:
        ws.protection.set_password("authorized-owner")
    wb.security = WorkbookProtection(lockStructure=structure, workbookPassword="owner")
    other = wb.create_sheet("Reference")
    other.merge_cells("A1:C1")
    other["A1"] = "Keep me"
    other.protection.sheet = True
    wb.properties.creator = "Auditor"
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def with_sharing(data, **attributes):
    result = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(result, "w") as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/workbook.xml":
                root = ET.fromstring(content)
                root.insert(0, ET.Element("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}fileSharing", attributes))
                content = ET.tostring(root)
            target.writestr(entry, content)
    return result.getvalue()


@pytest.mark.parametrize("recommended,structure", [(False, False), (True, False), (False, True)])
def test_separate_output_preserves_source_styles_and_protection(tmp_path, recommended, structure):
    data = workbook_bytes(structure=structure)
    if recommended:
        data = with_sharing(data, readOnlyRecommended="1", userName="Auditor")
    source, output = tmp_path / "source.xlsx", tmp_path / "updated.xlsx"
    source.write_bytes(data)
    result = process_workbook(source, output, "Sheet", 1, 2, 1)
    assert source.read_bytes() == data
    assert output.exists() and output != source
    assert result["source_info"]["read_only_recommended"] is recommended
    assert inspect_source(output)[1]["read_only_recommended"] is recommended
    before, after = load_workbook(BytesIO(data)), load_workbook(output)
    assert after.sheetnames == before.sheetnames
    assert [after.active.cell(r, 3).value for r in range(2, 6)] == [5, 12, 18, 28]
    assert after.active["C1"].value == "Rate(%)"
    assert after.active["A2"]._style == before.active["A2"]._style
    assert after.active.column_dimensions["A"].width == 25
    assert after.active.row_dimensions[2].height == 30
    assert after.security == before.security
    assert after["Reference"].protection == before["Reference"].protection
    assert str(after["Reference"].merged_cells) == "A1:C1"
    assert after.properties.creator == "Auditor"
    before.close()
    after.close()


def test_read_only_path(tmp_path):
    source, output = tmp_path / "readonly.xlsx", tmp_path / "output.xlsx"
    data = workbook_bytes()
    source.write_bytes(data)
    source.chmod(stat.S_IREAD)
    try:
        info = process_workbook(source, output, "Sheet", 1, 2, 1)["source_info"]
        if os.name == "nt":
            assert info["windows_read_only"] is True
        assert source.read_bytes() == data
        assert output.exists()
    finally:
        source.chmod(stat.S_IREAD | stat.S_IWRITE)


def test_uploaded_bytes_do_not_reopen_locked_original(tmp_path):
    # Once uploaded, processing needs no access to the original Excel file.
    data = workbook_bytes()
    assert inspect_source(data)[1]["windows_read_only"] is None
    source = tmp_path / "open-in-excel.xlsx"
    source.write_bytes(data)
    with source.open("rb"):
        process_workbook(data, tmp_path / "output.xlsx", "Sheet", 1, 2, 1)
    assert (tmp_path / "output.xlsx").exists()


@pytest.mark.parametrize("existing,recalculate", [(False, False), (True, True)])
def test_protected_selected_sheet_rejects_changes(tmp_path, existing, recalculate):
    with pytest.raises(PreservationError, match="worksheet is protected"):
        process_workbook(workbook_bytes(sheet_protected=True, existing=existing),
                         tmp_path / "output.xlsx", "Sheet", 1, 2, 1,
                         recalculate_existing=recalculate)
    assert not (tmp_path / "output.xlsx").exists()


def test_protected_sheet_can_keep_existing_rates_without_changes(tmp_path):
    data = workbook_bytes(sheet_protected=True, existing=True)
    output = tmp_path / "output.xlsx"
    process_workbook(data, output, "Sheet", 1, 2, 1)
    assert output.read_bytes() == data


@pytest.mark.parametrize("data,match", [
    (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1encrypted", "Encrypted/password-protected"),
    (b"invalid", "Invalid or unsupported"),
    (with_sharing(workbook_bytes(), reservationPassword="ABCD"), "password-to-modify"),
])
def test_security_restrictions_are_rejected(tmp_path, data, match):
    with pytest.raises(PreservationError, match=match):
        process_workbook(data, tmp_path / "output.xlsx", "Sheet", 1, 2, 1)
    assert not (tmp_path / "output.xlsx").exists()
