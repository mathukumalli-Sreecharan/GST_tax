from review_helpers import reviewed_post
from datetime import datetime, timedelta, timezone
from io import BytesIO
import json
from uuid import uuid4

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.auth import require_authenticated_user
from backend.main import app, RESULTS
from backend.maintenance import cleanup_expired_files


@pytest.fixture
def history(monkeypatch):
    user, other = str(uuid4()), str(uuid4())
    state = {'ready':True,'status':'approved','email_verified':True,'is_admin':False}
    rows, objects = {}, {}
    flags = {'upload_fails':False,'delete_fails':False,'reserve_fails':False,'complete_fails':False,'upload_ambiguous':False}
    current = [user]
    def rpc(name, values):
        identity=values.get('p_id')
        if name=='reserve_generated_file':
            if flags['reserve_fails']:raise HTTPException(503,'Database unavailable')
            item={'id':identity,'user_id':values['p_user'],'storage_path':values['p_user']+'/'+identity+'.xlsx',
                  'original_filename':values['p_original'],'output_filename':values['p_output'],
                  'file_size':values['p_size'],'processing_summary':values['p_summary'],
                  'verification_count':values['p_warnings'],'state':'uploading',
                  'created_at':datetime.now(timezone.utc).isoformat(),
                  'expires_at':(datetime.now(timezone.utc)+timedelta(days=30)).isoformat()}
            rows[identity]=item;return item
        if name=='claim_expired_generated_files':
            return [dict(row) for row in rows.values() if row.get('claimable')]
        if name=='finish_generated_file_delete':
            if identity in rows and rows[identity].get('cleanup_token')==values['p_token']:
                rows.pop(identity)
            return None
        if name=='list_generated_files':
            files=[row for row in rows.values() if row['user_id']==values['p_user'] and row['state']=='ready']
            return {'files':files,'total':len(files),'retention_days':30}
        row=rows.get(identity)
        if row is None or row['user_id']!=values.get('p_user'):
            raise HTTPException(404,'File not found')
        if name=='complete_generated_file':
            row['state']='ready'
            if flags['complete_fails']:raise HTTPException(503,'Response lost')
            return row
        if name in ('mark_generated_file_deleted','abandon_generated_file'):
            row['state']='deleting';row['cleanup_token']=str(uuid4());return row
        if name=='get_generated_file':
            if row['state']!='ready':raise HTTPException(404,'File not found')
            return row
        raise AssertionError(name)
    def storage(method,path,body=None,**kwargs):
        if method=='POST':
            if flags['upload_fails']:raise HTTPException(503,'Storage unavailable')
            objects[path.removeprefix('object/generated-excel-files/')]=body
            if flags['upload_ambiguous']:raise HTTPException(503,'Upload response lost')
            return b'{}'
        if method=='GET':return objects[path.removeprefix('object/authenticated/generated-excel-files/')]
        if method=='DELETE':
            if flags['delete_fails']:raise HTTPException(503,'Storage unavailable')
            for key in json.loads(body)['prefixes']:objects.pop(key,None)
            return b'[]'
        raise AssertionError(method)
    monkeypatch.setattr('backend.files.rpc',rpc)
    monkeypatch.setattr('backend.maintenance.rpc',rpc)
    monkeypatch.setattr('backend.files.storage_request',storage)
    monkeypatch.setattr('backend.approvals.approval_state',lambda _:dict(state))
    app.dependency_overrides[require_authenticated_user]=lambda:current[0]
    yield TestClient(app),rows,objects,flags,state,current,other
    app.dependency_overrides.clear()


def process(client, filename='private-audit.xlsx'):
    wb=Workbook();ws=wb.active;ws.append(['Invoice Value','Taxable Value'])
    for rate in (5,12,18,28):ws.append([100+rate,100])
    stream=BytesIO();wb.save(stream);wb.close()
    response=reviewed_post(client,'/api/excel/process',files={'file':(filename,stream.getvalue())},data={
        'worksheet':'Sheet','first_data_row':2,'invoice_col':1,'taxable_col':2})
    assert response.status_code==200,response.text
    return response.json()


def test_successful_save_is_private_and_downloads_across_sessions(history):
    client,rows,objects,_,_,_,_=history
    result=process(client)
    assert result['history']['saved'] is True
    identity=result['history']['id']
    assert rows[identity]['storage_path'].endswith(identity+'.xlsx')
    assert len(objects)==1
    # Lose all session-local data: persistent storage and metadata remain authoritative.
    RESULTS.clear()
    second=TestClient(app)
    listing=second.get('/api/files').json()
    assert listing['files'][0]['id']==identity and listing['retention_days']==30
    for path in (f'/api/files/{identity}/download',f'/api/excel/download/{identity}'):
        downloaded=second.get(path)
        assert downloaded.status_code==200
        assert downloaded.headers['cache-control']=='private, no-store'
        wb=load_workbook(BytesIO(downloaded.content))
        assert [wb.active.cell(r,3).value for r in range(2,6)]==[5,12,18,28]
        wb.close()


def test_ownership_on_every_history_download_and_delete(history):
    client,rows,_,_,_,current,other=history
    identity=process(client)['history']['id']
    current[0]=other
    assert client.get('/api/files').json()['files']==[]
    assert client.get(f'/api/files/{identity}/download').status_code==404
    assert client.delete(f'/api/files/{identity}').status_code==404
    assert identity in rows


@pytest.mark.parametrize('status',['pending','rejected','revoked'])
def test_history_requires_current_approval(history,status):
    client,_,_,_,state,_,_=history
    identity=process(client)['history']['id'];state['status']=status
    assert client.get('/api/files').status_code==403
    assert client.get(f'/api/files/{identity}/download').status_code==403
    assert client.delete(f'/api/files/{identity}').status_code==403
    assert client.get(f'/api/excel/download/{identity}').status_code==403


@pytest.mark.parametrize('failure',['upload_fails','reserve_fails','complete_fails','upload_ambiguous'])
def test_persistence_failure_keeps_session_download_without_claiming_saved(history,failure):
    client,rows,_,flags,_,_,_=history;flags[failure]=True
    result=process(client)
    assert result['history']['saved'] is False
    assert 'could not be confirmed' in result['history']['warning']
    assert client.get('/api/excel/download/'+result['file_id']).status_code==200
    assert client.get('/api/files').json()['files']==[]
    if rows:assert next(iter(rows.values()))['state']=='deleting'


def test_deletion_removes_object_and_prevents_session_copy_download(history):
    client,rows,objects,_,_,_,_=history
    result=process(client);identity=result['history']['id']
    assert client.delete('/api/files/'+identity).json()=={'deleted':True,'queued':False}
    assert rows=={} and objects=={}
    assert client.get('/api/files/'+identity+'/download').status_code==404
    assert client.get('/api/excel/download/'+result['file_id']).status_code==404


def test_failed_deletion_is_hidden_and_durable_cleanup_retries(history):
    client,rows,objects,flags,_,_,_=history
    identity=process(client)['history']['id'];flags['delete_fails']=True
    response=client.delete('/api/files/'+identity)
    assert response.status_code==202 and response.json()['queued'] is True
    assert client.get('/api/files').json()['files']==[]
    rows[identity]['claimable']=True
    assert cleanup_expired_files(max_batches=1)=={'removed':0,'deferred':1}
    assert identity in rows and objects
    flags['delete_fails']=False
    assert cleanup_expired_files(max_batches=1)=={'removed':1,'deferred':0}
    assert rows=={} and objects=={}


def test_history_query_validation(history):
    client,*_=history
    for query in ('sort=malicious','offset=-1','search='+'x'*201):
        assert client.get('/api/files?'+query).status_code==422


def test_history_retains_original_display_name_and_upload_is_removed(history):
    from backend.main import STORAGE
    client,rows,*_=history
    identity=process(client,'audit #1.xlsx')['history']['id']
    assert rows[identity]['original_filename']=='audit #1.xlsx'
    assert rows[identity]['output_filename']=='updated_audit _1.xlsx'
    assert list(STORAGE.glob('inspect_*'))==[]


@pytest.mark.parametrize('method,path', [('get','/api/files'),('get','/api/files/'+str(uuid4())+'/download'),('delete','/api/files/'+str(uuid4()))])
def test_history_without_authentication_is_denied(method,path):
    assert getattr(TestClient(app),method)(path).status_code==401
