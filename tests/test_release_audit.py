"""Phase 3 audit probes. All workbooks here are SYNTHETIC, not PRASUNA.

Phase 4 retains the audit cases and turns former defect characterizations into
positive regressions for their fixes.
"""
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile
import json
import subprocess
import sys

import pytest
from fastapi.testclient import TestClient
from fastapi import Request
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from backend.main import app
from backend.approvals import require_approved_user
from backend.reviews import connect, owned
from excel_processor import PreservationError, process_workbook
from gst_splits import calculate_split
from tests.test_reviews import save
from tests.test_splits import preview_confirm, values, fixture as split_fixture

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


@pytest.fixture
def client():
    app.dependency_overrides[require_approved_user] = lambda: 'audit-owner'
    yield TestClient(app)
    app.dependency_overrides.clear()


def synthetic_b2b():
    wb = Workbook(); ws = wb.active; ws.title = 'B2B'
    headers = {1:'GSTIN of supplier',2:'Supplier Name',3:'Invoice Number',4:'Invoice Date',
               6:'Invoice Value',9:'Taxable Value',10:'Original trailing detail',11:'IGST',
               12:'CGST',13:'SGST/UTGST',14:'Cess',16:'Invoice link',17:'Empty original text',18:'Last original column'}
    for col, label in headers.items():ws.cell(6,col,label)
    ws.merge_cells('A1:R1'); ws['A1']='SYNTHETIC AUDIT ONLY'
    transactions = [(348,300,48),(683,600,83),(None,None,None),(-118,-100,-18),(118,100,18)]
    for row,(invoice,taxable,igst) in enumerate(transactions,7):
        for col,value in {1:'GSTIN-AUDIT',2:'Synthetic Supplier',3:'AUDIT-'+str(row),4:'2026-10-25',
                          6:invoice,9:taxable,10:'KEEP-J-'+str(row),11:igst,
                          12:None if invoice is None else 0,13:None if invoice is None else 0,
                          14:None if invoice is None else 0,16:'Original external link',
                          17:'',18:'KEEP-LAST-'+str(row)}.items():ws.cell(row,col,value)
        ws.cell(row,16).hyperlink='https://example.org/audit'
        ws.cell(row,17).font=Font(bold=True,color='123456')
        ws.cell(row,17).fill=PatternFill('solid',fgColor='DCFCE7')
        ws.row_dimensions[row].height=22
    ws.merge_cells('E13:H13');ws['E13']='Notes: original footer'
    wb.create_sheet('Unrelated')['D8']='Original unrelated data'
    stream=BytesIO();wb.save(stream);wb.close()
    # openpyxl normally saves '' as an empty inline string without a text node.
    # Explicitly reproduce the source representation that reloads as data_type s, ''.
    output=BytesIO()
    with ZipFile(BytesIO(stream.getvalue())) as source,ZipFile(output,'w') as target:
        for entry in source.infolist():
            content=source.read(entry.filename)
            if entry.filename=='xl/worksheets/sheet1.xml':
                root=ET.fromstring(content)
                cell=next(cell for cell in root.iter(NS+'c') if cell.get('r')=='Q7')
                for child in list(cell):cell.remove(child)
                cell.set('t','inlineStr');ET.SubElement(ET.SubElement(cell,NS+'is'),NS+'t').text=''
                content=ET.tostring(root,encoding='utf-8')
            target.writestr(entry,content)
    return output.getvalue(),dict(worksheet='B2B',first_data_row=7,invoice_col=6,taxable_col=9)


def begin(client, raw, options):
    response=client.post('/api/excel/process',files={'file':('synthetic-audit.xlsx',raw)},data=options)
    assert response.status_code==200,response.text
    return response.json()


def test_synthetic_b2b_q7_journey_and_independent_reconciliation(client):
    raw,options=synthetic_b2b()
    original=load_workbook(BytesIO(raw))
    assert original['B2B']['Q7'].value=='' and original['B2B']['Q7'].data_type=='s'
    with patch('backend.main.process_workbook') as premature:
        inspected=client.post('/api/excel/inspect',files={'file':('synthetic-audit.xlsx',raw)})
        assert inspected.status_code==200
        review=begin(client,raw,options)
        premature.assert_not_called()
    assert [row['excel_row'] for row in review['rows']]==[7,8,9,10,11]
    assert review['rows'][2]['status']=='UNRESOLVED'  # all monetary cells blank
    assert review['rows'][3]['status']=='MATCHED'  # consistent signed credit uses the same rate rules
    path='/api/excel/reviews/'+review['review_id']
    assert client.post(path+'/generate').status_code==409
    assert client.post('/api/excel/generate').status_code==404
    assert client.post(path+'/rows/7',json={'version':review['version'],'choice':'WRONG','rates':[18,12],'slab_count':2,'split_confirmed':True}).status_code==422
    review=save(client,review,7,rates=[18,12]).json()
    review=preview_confirm(client,review,7)
    assert client.post(path+'/generate').status_code==409  # other required rows remain
    review=save(client,review,8,rates=[5,12,18]).json()
    missing=client.post(path+'/rows/8/split-preview',json={'version':review['version']}).json()
    assert missing['decisions']['8']['split_preview']['manual_required']
    assert client.post(path+'/generate').status_code==409
    review=preview_confirm(client,missing,8,[100,200,300])
    # A single selected rate preserves missing/negative original amounts; it does
    # not establish that signed mixed-rate splitting is supported.
    review=save(client,review,9,rates=[18]).json()
    review=save(client,review,10,rates=[18]).json()
    assigned=review['rows'][4]['rate']
    review=save(client,review,11,'CORRECT').json()
    assert review['rows'][4]['rate']==assigned
    persisted=client.get(path).json()
    assert persisted['decisions']['7']['split_confirmed'] and persisted['progress']['remaining']==0
    with connect() as db:assert owned(db,review['review_id'],'audit-owner')['source']==raw
    with patch('backend.main.save_generated_file',return_value={'saved':False}):
        response=client.post(path+'/generate')
    assert response.status_code==200,response.text
    result=response.json();download=client.get('/api/excel/download/'+result['file_id'])
    assert download.status_code==200
    with ZipFile(BytesIO(download.content)) as archive:assert archive.testzip() is None
    final=load_workbook(BytesIO(download.content));ws=final['B2B']
    assert final.sheetnames==original.sheetnames
    mapping=result['column_totals']['source_to_output_rows']
    assert mapping=={'7':[7,8],'8':[9,10,11],'9':[12],'10':[13],'11':[14]}
    assert ws['R7'].value is None and ws['R8'].value is None
    assert result['column_totals']['output']['row']==15
    splits={int(key):parts for key,parts in result['confirmed_splits'].items()}
    fields={'invoice':6,'taxable':9,'igst':11,'cgst':12,'sgst':13,'cess':14}
    def mapped_row(row):
        first=row+sum(len(parts)-1 for old,parts in splits.items() if old<row)
        return mapping.get(str(row),[first])
    # Independent comparison: uses source cells and API mapping, not validator internals.
    for old_ws in original:
        for old in list(old_ws._cells.values()):
            destinations=mapped_row(old.row) if old_ws.title=='B2B' else [old.row]
            col=old.column+(old_ws.title=='B2B' and old.column>=10)
            split_field=next((key for key,value in fields.items() if value==old.column),None)
            for i,row in enumerate(destinations):
                new=final[old_ws.title].cell(row,col)
                if old_ws.title=='B2B' and old.row in splits and split_field:
                    assert Decimal(str(new.value))==Decimal(splits[old.row][i][split_field])
                elif old.value=='' and old.data_type=='s':assert new.value in ('',None)
                else:assert new.value==old.value,(old_ws.title,old.coordinate,new.coordinate)
                assert new._style==old._style
                assert new.comment==old.comment
                if old.hyperlink:
                    assert new.hyperlink.target==old.hyperlink.target
                    assert new.hyperlink.ref==new.coordinate
    assert set(map(str,ws.merged_cells.ranges))=={'A1:S1','E16:H16'}
    expected={'invoice':Decimal('1031.00'),'taxable':Decimal('900.00'),'igst':Decimal('131.00'),
              'cgst':Decimal(0),'sgst':Decimal(0),'cess':Decimal(0)}
    for key,col in fields.items():
        source_total=sum(Decimal(str(original['B2B'].cell(row,col).value or 0)) for row in range(7,12))
        outcol=col+(col>=10)
        output_total=sum(Decimal(str(ws.cell(row,outcol).value or 0)) for row in range(7,15))
        assert source_total==output_total==expected[key]
        assert Decimal(str(ws.cell(15,outcol).value))==expected[key]
        assert Decimal(result['column_totals']['columns'][key]['total'])==expected[key]
    assert ws['J15'].value is None  # percentages never summed
    assert len([row for row in range(1,ws.max_row+1) if any(cell.value=='GRAND TOTAL' for cell in ws[row])])==1
    original.close();final.close()


def test_signed_mixed_credit_reconciles_after_phase4_fix():
    result=calculate_split(values('-348','-300','-48'),[18,12])
    assert [Decimal(row['taxable']) for row in result['rows']]==[-200,-100]
    assert Decimal(result['totals']['invoice'])==-348


def test_deletion_is_rejected_while_generation_owns_session(client):
    raw,options=synthetic_b2b();review=begin(client,raw,options)
    path='/api/excel/reviews/'+review['review_id']
    with connect() as db:
        db.execute("update reviews set state='generating' where id=?",(review['review_id'],))
    assert client.delete(path).status_code==409
    assert client.get(path).json()['state']=='generating'


def test_same_session_concurrent_edits_are_version_fenced(client):
    raw,options=synthetic_b2b();review=begin(client,raw,options)
    def edit(_):return save(client,review,7,'CORRECT').status_code
    with ThreadPoolExecutor(max_workers=2) as pool:statuses=list(pool.map(edit,range(2)))
    assert sorted(statuses)==[200,409]


def test_fractional_cent_total_uses_shared_rounding_after_phase4_fix(tmp_path):
    raw,options,header,_=split_fixture()
    wb=load_workbook(BytesIO(raw));ws=wb[options['worksheet']]
    ws.cell(header+3,options['invoice_col'],Decimal('118.005'))
    stream=BytesIO();wb.save(stream);wb.close()
    output=tmp_path/'fractional-cent.xlsx'
    parts=calculate_split(values(),[18,12])['rows']
    result=process_workbook(stream.getvalue(),output,options['worksheet'],options['invoice_col'],options['taxable_col'],header,
                         transaction_columns={'invoice_number':1,'invoice_date':2,'supplier_name':9},
                         rate_overrides={header+1:18},row_splits={header+1:parts})
    wb=load_workbook(output);col=options['invoice_col']+(options['invoice_col']>=options['taxable_col']+1)
    assert Decimal(result['column_totals']['columns']['invoice']['total'])==Decimal('1149.01')
    assert Decimal(str(wb[options['worksheet']].cell(result['column_totals']['output']['row'],col).value))==Decimal('1149.01')
    wb.close()


def test_note_in_supplier_column_is_structural_after_phase4_fix(client):
    raw,options=synthetic_b2b()
    wb=load_workbook(BytesIO(raw));ws=wb['B2B']
    ws.unmerge_cells('E13:H13');ws['E13']=None
    ws.merge_cells('A13:E13');ws['A13']='Notes: original footer'
    stream=BytesIO();wb.save(stream);wb.close()
    review=begin(client,stream.getvalue(),options)
    assert [row['excel_row'] for row in review['rows']]==[7,8,9,10,11]
    assert {'row':13,'reason':'trailing_note'} in review['excluded_rows']


def test_two_concurrent_users_have_isolated_sessions(client):
    def test_identity(request:Request):return request.headers['x-audit-owner']
    app.dependency_overrides[require_approved_user]=test_identity
    raw,options=synthetic_b2b()
    def create(owner):
        response=client.post('/api/excel/process',headers={'x-audit-owner':owner},files={'file':('synthetic.xlsx',raw)},data=options)
        assert response.status_code==200,response.text
        return owner,response.json()
    with ThreadPoolExecutor(max_workers=2) as pool:reviews=list(pool.map(create,['audit-a','audit-b']))
    for owner,review in reviews:
        path='/api/excel/reviews/'+review['review_id']
        assert client.get(path,headers={'x-audit-owner':owner}).status_code==200
        other='audit-b' if owner=='audit-a' else 'audit-a'
        assert client.get(path,headers={'x-audit-owner':other}).status_code==404
        assert client.post(path+'/generate',headers={'x-audit-owner':other}).status_code==404
        assert client.delete(path,headers={'x-audit-owner':other}).status_code==404


def test_local_sqlite_reopens_in_a_separate_process_when_disk_is_retained(client):
    raw,options=synthetic_b2b();review=begin(client,raw,options)
    review=save(client,review,7,'CORRECT').json()
    # This verifies local persistence only, not filesystem survival on Render.
    script=('import json; from backend.reviews import connect,owned,snapshot; '
            'db_context=connect(); db=db_context.__enter__(); '
            f'view=snapshot(owned(db,{review["review_id"]!r},"audit-owner")); '
            'print(json.dumps({"version":view["version"],"choice":view["decisions"]["7"]["choice"]})); '
            'db_context.__exit__(None,None,None)')
    result=subprocess.run([sys.executable,'-c',script],capture_output=True,text=True,check=True)
    assert json.loads(result.stdout)=={'version':review['version'],'choice':'CORRECT'}
