from review_helpers import reviewed_post
from pathlib import Path
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

import pytest
from backend.auth import require_authenticated_user
from backend.approvals import require_approved_user
from backend.main import app, configured_cors_origins

@pytest.fixture(autouse=True)
def authenticated_api():
    app.dependency_overrides[require_authenticated_user] = lambda: "test-user"
    app.dependency_overrides[require_approved_user] = lambda: "test-user"
    yield
    app.dependency_overrides.clear()



client = TestClient(app)


def test_cors_origins_are_trimmed_and_normalized():
    assert configured_cors_origins(" https://gst-tax-alpha.vercel.app/ , https://example.com ") == [
        "https://gst-tax-alpha.vercel.app", "https://example.com"
    ]


def test_inspect_and_process_regression_workbook():
    source = Path("KKENGIN APRIL-26.xlsx")
    if not source.exists():
        pytest.skip("Optional real workbook fixture KKENGIN APRIL-26.xlsx is unavailable")
    data = source.read_bytes()
    inspected = client.post("/api/excel/inspect", files={"file": (source.name, data,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert inspected.status_code == 200
    body = inspected.json()
    sheet = body["sheets"][0]
    assert sheet["name"] == "B2B"
    assert sheet["detected"]["invoice_col"] == 6
    assert sheet["detected"]["taxable_col"] == 9

    processed = reviewed_post(client,"/api/excel/process", files={"file": (source.name, data,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, data={
            "worksheet": "B2B", "first_data_row": "7", "invoice_col": "6", "taxable_col": "9",
            "invoice_number_col": "3", "invoice_header_row": "6", "taxable_header_row": "5",
        })
    assert processed.status_code == 200, processed.text
    result = processed.json()
    assert result["rate_column"] == {"letter": "J", "index": 10}
    assert result["summary"]["total"] > 0
    assert all(row["status"] != "MATCHED" for row in result["verification"])

    output = client.get(f"/api/excel/download/{result['file_id']}")
    report = client.get(f"/api/excel/report/{result['file_id']}")
    assert output.status_code == 200 and output.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument")
    assert "filename*=utf-8''updated_KKENGIN%20APRIL-26.xlsx" in output.headers["content-disposition"]
    assert report.status_code == 200 and b"Excel row" in report.content


def test_download_rejects_unknown_file_id():
    response = client.get("/api/excel/download/not-a-real-id")
    assert response.status_code == 404


@pytest.mark.parametrize("recommended", [False, True])
def test_inspection_only_reports_detectable_read_only_metadata(recommended):
    from test_read_only import workbook_bytes, with_sharing
    data = workbook_bytes()
    if recommended:
        data = with_sharing(data, readOnlyRecommended="1")
    response = client.post("/api/excel/inspect", files={"file": ("source.xlsx", data)})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["source_info"] == {"windows_read_only": None,
        "read_only_recommended": recommended, "read_only_detected": recommended}
    assert body["protection"]["protected_worksheets"] == ["Reference"]


@pytest.mark.parametrize("endpoint", ["inspect", "process"])
def test_encrypted_upload_has_specific_error(endpoint):
    response = client.post(f"/api/excel/{endpoint}",
        files={"file": ("encrypted.xlsx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1encrypted")},
        data={"worksheet": "Sheet", "first_data_row": "2", "invoice_col": "1", "taxable_col": "2"})
    assert response.status_code == 422
    assert "Encrypted/password-protected" in response.json()["detail"]


def test_api_reports_all_four_rates_and_exports_rate_column():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Invoices"
    sheet.append(["Invoice Value", "Taxable Value"])
    for rate in (5, 12, 18, 28):
        sheet.append([100 + rate, 100])
    stream = BytesIO()
    workbook.save(stream)

    response = reviewed_post(client,"/api/excel/process", files={"file": ("four-rates.xlsx", stream.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, data={
            "worksheet": "Invoices", "first_data_row": "2", "invoice_col": "1", "taxable_col": "2",
            "invoice_header_row": "1", "taxable_header_row": "1",
        })
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["summary"]["rate_counts"] == {"5": 1, "12": 1, "18": 1, "28": 1}
    assert [result["summary"][key] for key in
            ("five_percent", "twelve_percent", "eighteen_percent", "twenty_eight_percent")] == [1, 1, 1, 1]
    assert all(row["status"] == "MATCHED" for row in result["rows"])
    assert all(row["expected_12"] and row["difference_28"] for row in result["rows"])

    output = client.get(f"/api/excel/download/{result['file_id']}")
    assert output.status_code == 200
    exported = load_workbook(BytesIO(output.content), data_only=False)
    exported_sheet = exported["Invoices"]
    assert [exported_sheet.cell(1, column).value for column in (1, 2, 3)] == [
        "Invoice Value", "Taxable Value", "Rate(%)"
    ]
    assert [exported_sheet.cell(row, 3).value for row in range(2, 6)] == [5, 12, 18, 28]
    assert [exported_sheet.cell(row, 1).value for row in range(2, 6)] == [105, 112, 118, 128]
    assert [exported_sheet.cell(row, 2).value for row in range(2, 6)] == [100, 100, 100, 100]
    exported.close()


def test_api_rate_summary_includes_zero_counts():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Invoice Value", "Taxable Value"])
    sheet.append([105, 100])
    stream = BytesIO()
    workbook.save(stream)
    response = reviewed_post(client,"/api/excel/process", files={"file": ("one-rate.xlsx", stream.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, data={
            "worksheet": "Sheet", "first_data_row": "2", "invoice_col": "1", "taxable_col": "2",
            "invoice_header_row": "1", "taxable_header_row": "1",
        })
    assert response.status_code == 200, response.text
    assert response.json()["summary"]["rate_counts"] == {"5": 1, "12": 0, "18": 0, "28": 0}


def test_api_rejects_non_xlsx_upload():
    response = client.post("/api/excel/inspect", files={"file": ("book.xls", b"not excel", "application/octet-stream")})
    assert response.status_code == 400


def test_existing_rate_column_requires_confirmation_and_is_reused():
    source = Path("updated_KK MAY.xlsx")
    if not source.exists():
        pytest.skip("Optional real workbook fixture updated_KK MAY.xlsx is unavailable")
    content = source.read_bytes()
    upload = {"file": (source.name, content,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    base_fields = {"worksheet": "B2B", "first_data_row": "7", "invoice_col": "6", "taxable_col": "9",
                   "invoice_number_col": "3", "invoice_header_row": "6", "taxable_header_row": "5"}

    refused = reviewed_post(client,"/api/excel/process", files=upload,
                          data={**base_fields, "recalculate_existing": "true", "confirm_overwrite": "false"})
    assert refused.status_code == 400
    assert "Confirm overwrite" in refused.json()["detail"]

    kept = reviewed_post(client,"/api/excel/process", files=upload, data=base_fields)
    assert kept.status_code == 200, kept.text
    keep_result = kept.json()
    assert keep_result["rate_column"] == {"letter": "J", "index": 10}
    kept_download = client.get(f"/api/excel/download/{keep_result['file_id']}")
    assert kept_download.content == content

    confirmed = reviewed_post(client,"/api/excel/process", files=upload,
                            data={**base_fields, "recalculate_existing": "true", "confirm_overwrite": "true"})
    assert confirmed.status_code == 200, confirmed.text
    confirmed_result = confirmed.json()
    workbook_bytes = client.get(f"/api/excel/download/{confirmed_result['file_id']}").content
    from io import BytesIO
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)
    sheet = workbook["B2B"]
    assert sheet.max_column == 22
    assert sum(1 for row in sheet.iter_rows(min_row=1, max_row=40) for cell in row
               if str(cell.value or "").strip().lower() == "rate(%)") == 1
    workbook.close()
