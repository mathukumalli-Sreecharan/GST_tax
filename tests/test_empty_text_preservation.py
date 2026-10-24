"""Empty text normalization must never conceal nonempty source data loss."""
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from unittest.mock import patch

import pytest
from openpyxl import Workbook,load_workbook
from openpyxl.utils import get_column_letter

from excel_processor import process_workbook, PreservationError
from workbook_preservation import _is_empty_text_normalization
from test_numeric_preservation import numeric_fixture, NS


def empty_text_source(data, coordinate):
    output=BytesIO()
    with ZipFile(BytesIO(data)) as source,ZipFile(output,'w') as target:
        for item in source.infolist():
            content=source.read(item.filename)
            if item.filename=='xl/worksheets/sheet1.xml':
                root=ET.fromstring(content)
                cell=next(cell for cell in root.iter(NS+'c') if cell.get('r')==coordinate)
                for child in list(cell):cell.remove(child)
                cell.set('t','inlineStr')
                ET.SubElement(ET.SubElement(cell,NS+'is'),NS+'t').text=''
                content=ET.tostring(root,encoding='utf-8')
            target.writestr(item,content)
    return output.getvalue()


@pytest.mark.parametrize('seed',range(6))
def test_explicit_empty_text_survives_dynamic_mapping_as_blank(tmp_path,seed):
    data,title,header,cols,row,tail=numeric_fixture(seed,'12345678901234567')
    coordinate=f'{get_column_letter(tail)}{row}'
    # Remove the hyperlink: Excel/OpenPyXL fills blank linked cells with link text.
    model=load_workbook(BytesIO(data));model[title][coordinate].hyperlink=None
    stream=BytesIO();model.save(stream);model.close()
    data=empty_text_source(stream.getvalue(),coordinate)
    original=load_workbook(BytesIO(data));old=original[title][coordinate]
    assert old.value=='' and old.data_type=='s'
    shifted=cols['taxable']+1
    ordinary=load_workbook(BytesIO(data));ordinary[title].insert_cols(shifted)
    stream=BytesIO();ordinary.save(stream)
    roundtrip=load_workbook(BytesIO(stream.getvalue()))
    assert roundtrip[title].cell(row,tail+1).value is None
    output=tmp_path/'normalized.xlsx'
    result=process_workbook(data,output,title,cols['invoice'],cols['taxable'],header)
    actual=load_workbook(output);new=actual[title].cell(row,tail+1)
    assert new.value is None
    assert new._style==old._style
    assert result['column_totals']['output']['row']==row+1
    assert result['column_totals']['columns']['invoice']['total']=='118.00'
    repeated=tmp_path/'repeated.xlsx'
    process_workbook(output,repeated,title,cols['invoice']+(cols['invoice']>=shifted),cols['taxable'],header)
    replay=load_workbook(repeated)
    assert replay[title].max_row==actual[title].max_row
    for book in (original,ordinary,roundtrip,actual,replay):book.close()


@pytest.mark.parametrize('value,kind,new_value,new_kind,allowed',[
    ('','s',None,'inlineStr',True),('','s',None,'n',True),
    (' ','s',None,'n',False),('\t','s',None,'n',False),
    ('0','s',None,'n',False),('original text','s',None,'n',False),
    (0,'n',None,'n',False),(12345678901234567,'n',None,'n',False),
    ('=A1+B1','f',None,'n',False),('','f',None,'n',False),
    ('','s',None,'f',False),('','s',0,'n',False),
])
def test_normalization_is_limited_to_empty_text(value,kind,new_value,new_kind,allowed):
    wb=Workbook();old=wb.active['A1'];new=wb.active['B1']
    old.value=value;old.data_type=kind;new.value=new_value;new.data_type=new_kind
    assert _is_empty_text_normalization(old,new) is allowed


def test_real_nonempty_text_loss_still_refuses_export(tmp_path):
    data,title,header,cols,row,tail=numeric_fixture(70,'12345678901234567')
    model=load_workbook(BytesIO(data));model[title].cell(row,tail,'Original nonempty text')
    stream=BytesIO();model.save(stream)
    save=Workbook.save
    def damage(book,path):
        book[title].cell(row,tail+1).value=None
        save(book,path)
    output=tmp_path/'damaged.xlsx'
    with patch.object(Workbook,'save',damage):
        with pytest.raises(PreservationError,match='original value'):
            process_workbook(stream.getvalue(),output,title,cols['invoice'],cols['taxable'],header)
    assert not output.exists()
    model.close()


def test_empty_text_normalization_does_not_excuse_lost_style(tmp_path):
    data,title,header,cols,row,tail=numeric_fixture(71,'12345678901234567')
    coordinate=f'{get_column_letter(tail)}{row}'
    model=load_workbook(BytesIO(data));model[title][coordinate].hyperlink=None
    stream=BytesIO();model.save(stream);model.close()
    data=empty_text_source(stream.getvalue(),coordinate)
    save=Workbook.save
    def damage(book,path):
        book[title].cell(row,tail+1)._style=None
        save(book,path)
    output=tmp_path/'lost-style.xlsx'
    with patch.object(Workbook,'save',damage):
        with pytest.raises(PreservationError,match='style at'):
            process_workbook(data,output,title,cols['invoice'],cols['taxable'],header)
    assert not output.exists()
