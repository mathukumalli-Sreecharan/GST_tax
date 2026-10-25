"""Verification-first API, private review storage, and generation gate."""
from io import BytesIO
import json
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.main import app, RESULTS
from backend.approvals import require_approved_user
from backend.reviews import connect


def source(layout=0,count=1):
    wb=Workbook();ws=wb.active;ws.title='Transactions '+str(layout)
    fields=['invoice_number','supplier_gstin','supplier_name','invoice_date','invoice','taxable','igst','cgst','sgst','cess']
    if layout:fields=list(reversed(fields))
    labels={'invoice_number':'Invoice No.','supplier_gstin':'GSTIN of supplier','supplier_name':'Supplier Name','invoice_date':'Invoice Date','invoice':'Invoice Value','taxable':'Taxable Value','igst':'Integrated Tax','cgst':'Central Tax','sgst':'State/UT Tax','cess':'Cess'}
    header=2+layout
    for col,key in enumerate(fields,1):ws.cell(header,col,labels[key])
    examples=[(118,100,18),(None,None,None),(180,150,30),(125,100,18),(None,100,None)]
    for i in range(count):
        inv,tax,igst=examples[i%len(examples)]
        data={'invoice_number':'INV-'+str(i),'supplier_gstin':'TEST-GSTIN','supplier_name':'Supplier','invoice_date':'2026-10-25','invoice':inv,'taxable':tax,'igst':igst,'cgst':None if inv is None and tax is None else 0,'sgst':None if inv is None and tax is None else 0,'cess':None if inv is None and tax is None else 0}
        for col,key in enumerate(fields,1):ws.cell(header+1+i,col,data[key])
    stream=BytesIO();wb.save(stream)
    return stream.getvalue(),{'worksheet':ws.title,'first_data_row':header+1,'invoice_col':fields.index('invoice')+1,'taxable_col':fields.index('taxable')+1},header


@pytest.fixture
def client():
    app.dependency_overrides[require_approved_user]=lambda:'review-owner'
    yield TestClient(app)
    app.dependency_overrides.clear()


def begin(client,layout=0,count=5):
    data,fields,header=source(layout,count)
    response=client.post('/api/excel/process',files={'file':('synthetic.xlsx',data)},data=fields)
    assert response.status_code==200,response.text
    return response.json(),data,fields,header


def save(client,review,row,choice='WRONG',rates=None):
    body={'version':review['version'],'choice':choice}
    if choice=='WRONG':body.update(slab_count=len(rates or [5,12]),rates=rates or [5,12])
    return client.post(f"/api/excel/reviews/{review['review_id']}/rows/{row}",json=body)


@pytest.mark.parametrize('layout',[0,1])
def test_all_transactions_including_empty_money_reach_review_and_no_excel_is_generated(client,layout):
    with patch('backend.main.process_workbook') as export,patch('backend.main.save_generated_file') as history:
        review,raw,_,header=begin(client,layout)
        export.assert_not_called();history.assert_not_called()
    assert len(review['rows'])==5
    assert [row['excel_row'] for row in review['rows']]==list(range(header+1,header+6))
    assert review['rows'][0]['status']=='MATCHED'
    assert review['rows'][1]['status']=='UNRESOLVED'
    assert review['rows'][1]['invoice_number']=='INV-1'
    assert review['rows'][1]['supplier_gstin']=='TEST-GSTIN'
    assert review['rows'][2]['status']=='REVIEW_REQUIRED' # mixed ratio, engine only proposes a rate
    assert review['rows'][3]['status']=='REVIEW_REQUIRED' # inconsistent components
    assert review['rows'][4]['status']=='UNRESOLVED'
    assert review['progress']=={'total':5,'matched':1,'requiring_review':4,'verified':0,'remaining':4}
    assert 'file_id' not in review
    with connect() as db:
        assert db.execute('select source from reviews where id=?',(review['review_id'],)).fetchone()[0]==raw


def test_gate_decisions_reload_edits_and_slab_validation(client):
    review,raw,_,_=begin(client)
    path='/api/excel/reviews/'+review['review_id']
    assert client.post(path+'/generate').status_code==409
    assert client.get('/api/excel/download/'+review['review_id']).status_code!=200
    missing=review['rows'][1]['excel_row']
    assert save(client,review,missing,'CORRECT').status_code==422
    row=review['rows'][2]['excel_row']
    assert client.post(f'{path}/rows/{row}',json={'version':1,'choice':'WRONG','slab_count':2,'rates':[18]}).status_code==422
    response=save(client,review,row,'CORRECT');assert response.status_code==200
    updated=response.json()
    assert updated['rows'][2]['rate']==review['rows'][2]['rate']
    assert updated['progress']['verified']==1
    assert save(client,review,row).status_code==409 # stale version
    response=save(client,updated,row,rates=[18,28]);assert response.status_code==200
    persisted=client.get(path).json()
    assert persisted['decisions'][str(row)]['rates']==['18','28']
    assert persisted['version']==3
    assert client.get('/api/excel/reviews').json()['reviews'][0]['review_id']==review['review_id']
    for transaction in persisted['rows']:
        if transaction['status']!='MATCHED' and str(transaction['excel_row']) not in persisted['decisions']:
            response=save(client,persisted,transaction['excel_row'],rates=[18]);assert response.status_code==200
            persisted=response.json()
    # Multi-rate decisions additionally require a financially valid preview and confirmation.
    response=client.post(f'{path}/rows/{row}/split-preview',json={'version':persisted['version']})
    assert response.status_code==200,response.text
    persisted=response.json()
    preview=persisted['decisions'][str(row)]['split_preview']
    response=client.post(f'{path}/rows/{row}/split-confirm',json={'version':persisted['version'],'preview_id':preview['preview_id']})
    assert response.status_code==200,response.text
    persisted=response.json()
    assert persisted['progress']['remaining']==0
    with patch('backend.main.save_generated_file',return_value={'saved':False}):
        result=client.post(path+'/generate')
    assert result.status_code==200,result.text
    output=result.json()
    downloaded=client.get('/api/excel/download/'+output['file_id'])
    assert downloaded.status_code==200
    wb=load_workbook(BytesIO(downloaded.content));ws=wb[output['worksheet']]
    assert ws.cell(row,output['rate_column']['index']).value==18
    assert ws.cell(row+1,output['rate_column']['index']).value==28
    assert ws.cell(max(item['excel_row'] for item in persisted['rows'])+2,1).value=='GRAND TOTAL'
    assert client.post(path+'/generate').json()['file_id']==output['file_id']
    assert save(client,persisted,row).status_code==409
    with connect() as db:assert db.execute('select source from reviews where id=?',(review['review_id'],)).fetchone()[0]==raw
    wb.close()


def test_private_access_and_expiry(client):
    review,_,_,_=begin(client,count=1)
    path='/api/excel/reviews/'+review['review_id']
    app.dependency_overrides[require_approved_user]=lambda:'other-user'
    assert client.get(path).status_code==404
    assert client.post(path+'/generate').status_code==404
    assert client.delete(path).status_code==404
    assert client.get('/api/excel/reviews').json()['reviews']==[]
    app.dependency_overrides[require_approved_user]=lambda:'review-owner'
    with connect() as db:db.execute('update reviews set expires=0 where id=?',(review['review_id'],))
    assert client.get(path).status_code==404


def test_generation_failure_preserves_review_and_can_retry(client):
    review,_,_,_=begin(client,count=1)
    path='/api/excel/reviews/'+review['review_id']
    with patch('backend.main.generate_excel',side_effect=ValueError('test generation failure')):
        with pytest.raises(ValueError):client.post(path+'/generate')
    assert client.get(path).json()['state']=='open'


def test_review_endpoints_require_login():
    app.dependency_overrides.clear()
    client=TestClient(app)
    assert client.get('/api/excel/reviews').status_code==401
    assert client.get('/api/excel/reviews/private').status_code==401
    assert client.post('/api/excel/reviews/private/generate').status_code==401


@pytest.mark.parametrize('choice', ['CORRECT', 'WRONG'])
def test_existing_rate_confirmation_is_unchanged_and_correction_is_explicit(client, choice):
    raw, fields, header = source(count=1)
    wb = load_workbook(BytesIO(raw)); ws = wb.active
    col = ws.max_column + 1
    ws.cell(header, col, 'Rate (%)'); ws.cell(header + 1, col, 12)
    stream = BytesIO(); wb.save(stream); wb.close()
    review = client.post('/api/excel/process', files={'file': ('existing.xlsx', stream.getvalue())}, data=fields).json()
    row = header + 1
    assert review['rows'][0]['rate'] == 12
    assert review['rows'][0]['status'] == 'REVIEW_REQUIRED'
    response = save(client, review, row, choice, rates=[5])
    assert response.status_code == 200, response.text
    with patch('backend.main.save_generated_file', return_value={'saved': False}):
        response = client.post('/api/excel/reviews/' + review['review_id'] + '/generate')
    assert response.status_code == 200, response.text
    result = response.json()
    expected = 12 if choice == 'CORRECT' else 5
    assert result['rows'][0]['rate'] == expected
    output = client.get('/api/excel/download/' + result['file_id'])
    wb = load_workbook(BytesIO(output.content))
    assert wb.active.cell(row, col).value == expected
    wb.close()
