"""Conservative monetary aggregation, independent of GST rate classification."""
from copy import copy
from decimal import Decimal, ROUND_HALF_UP, localcontext
import re
import hashlib

from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import quote_sheetname
from openpyxl.styles import Font, PatternFill, Side
from openpyxl.utils import get_column_letter, column_index_from_string

FIELDS = {'invoice': 'Invoice Value', 'taxable': 'Taxable Value', 'igst': 'IGST',
          'cgst': 'CGST', 'sgst': 'SGST/UTGST', 'cess': 'Cess'}
CURRENCY = '"₹" #,##0.00;[Red]-"₹" #,##0.00'


def norm(value):
    return re.sub(r'[^a-z0-9]+', ' ', str(value or '').lower()).strip()


def header_kind(value):
    text = norm(value)
    words = set(text.split())
    compact = text.replace(' ', '')
    if compact in {'igst', 'igstamount', 'igstvalue', 'igstrs', 'igstinr'}: return 'igst'
    if compact in {'cgst', 'cgstamount', 'cgstvalue', 'cgstrs', 'cgstinr'}: return 'cgst'
    if compact in {'sgst', 'utgst', 'sgstutgst', 'sgstamount', 'utgstamount'}: return 'sgst'
    acronyms = words & {'igst', 'cgst', 'sgst', 'utgst'}
    if len(acronyms) > 1 and not acronyms <= {'sgst', 'utgst'}:
        return None
    if ('invoice' in words or 'inv' in words) and words & {'value', 'amount', 'amt'}:
        return 'invoice'
    if 'taxable' in words and words & {'value', 'amount', 'amt'}:
        return 'taxable'
    if ('igst' in words or ('integrated' in words and 'tax' in words)) and not words & {'rate', 'percent'}:
        return 'igst'
    if ('cgst' in words or ('central' in words and 'tax' in words)) and not words & {'rate', 'percent'}:
        return 'cgst'
    if (words & {'sgst', 'utgst'} or compact.startswith('sgstutgst') or
            (words & {'state', 'ut', 'union'} and 'tax' in words)) and not words & {'rate', 'percent'}:
        return 'sgst'
    if 'cess' in words and not words & {'rate', 'percent'}:
        return 'cess'
    return None


def header_at(ws, cell):
    """Recognize vertically split labels/units without reading transaction text as a header."""
    key = header_kind(cell.value)
    if key:
        return key
    fragment = norm(cell.value)
    if fragment not in {'value', 'amount', 'amt', 'tax', 'inr', 'rs', 'rupees'} and str(cell.value or '').strip() != '\u20b9':
        return None
    parts = [ws.cell(row, cell.column).value for row in range(max(1, cell.row - 2), cell.row + 1)]
    if not all(value is None or isinstance(value, str) for value in parts):
        return None
    return header_kind(' '.join(str(value or '') for value in parts))


def detect_monetary_columns(ws, last_header_row=None):
    candidates = {key: set() for key in FIELDS}
    for row in ws.iter_rows(max_row=min(ws.max_row, last_header_row) if last_header_row is not None else ws.max_row):
        for cell in row:
            key = header_at(ws, cell)
            if key and '%' not in str(cell.value):
                candidates[key].add(cell.column)
    columns = {key: next(iter(cols)) if len(cols) == 1 else None for key, cols in candidates.items()}
    warnings = [f'{FIELDS[key]} column is {"ambiguous" if cols else "missing"}; select a column or review.'
                for key, cols in candidates.items() if len(cols) != 1 and key != 'cess']
    if len(candidates['cess']) > 1:
        warnings.append('Cess column is ambiguous; select a column or review.')
    return {'columns': columns, 'candidates': {k: sorted(v) for k, v in candidates.items()}, 'warnings': warnings}


def marker_name(ws, kind):
    return '_GST_LEDGER_' + kind + '_' + hashlib.sha256(ws.title.encode()).hexdigest()[:16]


def marker_destination(ws, kind):
    name = ws.parent.defined_names.get(marker_name(ws, kind))
    if name is None:
        return None
    try:
        sheet, reference = next(iter(name.destinations))
        return sheet.replace("''", "'"), reference
    except (ValueError, StopIteration, TypeError):
        return None


def generated_row(ws, row):
    destination = marker_destination(ws, 'TOTAL')
    if not destination or destination[0] != ws.title:
        return False
    from openpyxl.utils.cell import coordinate_from_string
    try:
        col, number = coordinate_from_string(destination[1].replace('$', ''))
        return number == row and ws.cell(row, column_index_from_string(col)).value == 'GRAND TOTAL'
    except ValueError:
        return False


def excluded_reason(ws, row, columns):
    if generated_row(ws, row):
        return 'generated_total'
    # Recognize labels, not substrings in descriptions or invoice identifiers.
    for cell in ws[row]:
        if isinstance(cell.value, str) and re.fullmatch(
                r'(?:grand\s*total|sub\s*total|total)(?:\s+(?:amount|value|inr|rs|invoice value|taxable value|igst|cgst|sgst|cess))?', norm(cell.value)):
            return 'existing_total'
    if any(isinstance(cell.value, str) and re.match(r'^(?:notes?|remarks?|prepared by|disclaimer)\b', norm(cell.value)) for cell in ws[row]) and not any(
            columns.get(key) and monetary_value(ws.cell(row, columns[key]))[1] == 'numeric' for key in ('invoice', 'taxable')):
        return 'trailing_note'
    selected = [ws.cell(row, col) for col in columns.values() if col]
    header_matches = sum(header_at(ws, cell) is not None or norm(cell.value) in {
        'invoice', 'taxable', 'integrated', 'central', 'state', 'state ut', 'cess',
        'value', 'amount', 'inr', 'rs', 'rupees'} for cell in selected)
    if header_matches >= 2 and not any(isinstance(cell.value, (int, float, Decimal)) for cell in selected):
        return 'repeated_header'
    if not any(ws.cell(row, col).value not in (None, '') for col in columns.values() if col):
        return 'blank'
    return None


def monetary_value(cell):
    if cell.value is None or cell.value == '':
        return Decimal(0), 'blank'
    if (isinstance(cell.value, bool) or not isinstance(cell.value, (int, float, Decimal))
            or cell.is_date or '%' in cell.number_format or cell.data_type in {'f', 'e'}):
        return None, 'invalid'
    value = Decimal(str(cell.value))
    return (value, 'numeric') if value.is_finite() else (None, 'invalid')


def aggregate(ws, first_data_row, invoice_col, taxable_col, overrides=None):
    if not 2 <= first_data_row <= ws.max_row + 1:
        raise ValueError('First data row must follow the headers and lie within the worksheet.')
    detected = detect_monetary_columns(ws, first_data_row - 1)
    columns = {**detected['columns'], 'invoice': invoice_col, 'taxable': taxable_col}
    for key, col in (overrides or {}).items():
        if key not in {'igst', 'cgst', 'sgst', 'cess'}:
            raise ValueError('Only IGST, CGST, SGST/UTGST and Cess mappings may be overridden.')
        if col is not None and (isinstance(col, bool) or not isinstance(col, int) or not 1 <= col <= ws.max_column):
            raise ValueError('Monetary column selection is outside the worksheet.')
        columns[key] = col
    chosen = [col for col in columns.values() if col is not None]
    if len(set(chosen)) != len(chosen) or any(not 1 <= col <= ws.max_column for col in chosen):
        raise ValueError('Each monetary field must use a different valid source column.')
    for col in chosen:
        for row in range(1, first_data_row):
            value = ws.cell(row, col).value
            text = norm(value)
            if '%' in str(value or '') or set(text.split()) & {'rate', 'percent', 'percentage'}:
                raise ValueError('Rate (%) or percentage columns cannot be monetary columns.')
            if text in {'invoice number', 'invoice no', 'inv no', 'inv number', 'bill no',
                        'bill number', 'document number', 'document no', 'gstin'}:
                raise ValueError('Invoice identifiers cannot be monetary columns.')
    warnings = [f'{FIELDS[key]} column is unavailable; total is not calculated.' for key, col in columns.items()
                if col is None and (key != 'cess' or detected['candidates']['cess'])]
    identifiers = set()
    for header in ws.iter_rows(max_row=first_data_row - 1):
        for cell in header:
            if norm(cell.value) in {'invoice number', 'invoice no', 'inv no', 'document number', 'document no'}:
                identifiers.add(cell.column)
    boundary_conflicts = []
    included, excluded, issues = [], [], []
    sums = {key: Decimal(0) for key in FIELDS}
    blanks = {key: 0 for key in FIELDS}
    invalid = {key: 0 for key in FIELDS}
    with localcontext() as ctx:
        ctx.prec = 50
        for row in range(first_data_row, ws.max_row + 1):
            reason = excluded_reason(ws, row, columns)
            if reason:
                if reason == 'trailing_note' and any(monetary_value(ws.cell(row, col))[1] == 'numeric' for col in columns.values() if col):
                    boundary_conflicts.append(row)
                if reason != 'blank':
                    excluded.append({'row': row, 'reason': reason})
                continue
            if not any(monetary_value(ws.cell(row, col))[1] == 'numeric' for col in columns.values() if col) and not any(
                    ws.cell(row, col).value not in (None, '') for col in identifiers):
                boundary_conflicts.append(row)
            included.append(row)
            for key, col in columns.items():
                if col is None:
                    continue
                value, state = monetary_value(ws.cell(row, col))
                if state == 'invalid':
                    invalid[key] += 1
                    issues.append({'row': row, 'field': key, 'column': col, 'message': 'Invalid monetary value; excluded from this column total.'})
                else:
                    sums[key] += value
                    blanks[key] += state == 'blank'
        totals = {key: {'label': FIELDS[key], 'source_column': col,
                        'total': str(sums[key].quantize(Decimal('.01'), rounding=ROUND_HALF_UP)) if col else None,
                        'blank_count': blanks[key], 'invalid_count': invalid[key]}
                  for key, col in columns.items()}
    if not included:
        warnings.append('No transaction rows were included; review the worksheet and first data row.')
    if issues:
        warnings.append('Invalid monetary cells were excluded from their column totals. Review the listed cells before filing.')
    return {'worksheet': ws.title, 'columns': totals, 'included_rows': included, 'excluded_rows': excluded,
            'issues': issues, 'warnings': warnings, 'output': None, 'boundary_conflicts': boundary_conflicts}


class TotalsConflict(ValueError):
    """The mandatory original-sheet total cannot be written without guessing or damage."""


def write_totals(wb, ws, totals, shift_column=None):
    """Write immediately after transactions, or refuse an occupied/unsafe destination."""
    def mapped(col):
        return col + 1 if shift_column and col >= shift_column else col
    columns = {key: mapped(item['source_column']) for key, item in totals['columns'].items() if item['source_column']}
    if not totals['included_rows']:
        raise TotalsConflict('No transaction boundary could be identified. Review the first data row and monetary columns.')
    if ws.protection.sheet:
        raise TotalsConflict('The selected worksheet is protected; a GRAND TOTAL row cannot be written without bypassing protection.')
    if totals.get('boundary_conflicts'):
        raise TotalsConflict('Transaction boundaries are ambiguous at row(s) ' + ', '.join(map(str, totals['boundary_conflicts'][:10])) + '. Review invalid-only monetary rows or unrelated content before exporting.')
    last = max(totals['included_rows'])
    if any(item['reason'] == 'trailing_note' and item['row'] < last for item in totals['excluded_rows']):
        raise TotalsConflict('Monetary content follows a notes section; transaction boundaries require verification.')
    row = last + 1
    if row > 1048576:
        raise TotalsConflict('There is no Excel row available after the last transaction.')
    generated = [r for r in range(1, ws.max_row + 1) if generated_row(ws, r)]
    if generated and generated != [row]:
        raise TotalsConflict('The generated GRAND TOTAL row is no longer immediately after the transactions. Review the layout before processing.')
    if any(area.min_row <= row <= area.max_row for area in ws.merged_cells.ranges):
        raise TotalsConflict(f'GRAND TOTAL destination row {row} intersects merged cells. No content was overwritten.')
    existing_label = marker_destination(ws, 'TOTAL') if generated else None
    if existing_label:
        from openpyxl.utils.cell import coordinate_from_string
        label_col = column_index_from_string(coordinate_from_string(existing_label[1].replace('$', ''))[0])
        if label_col in columns.values():
            raise TotalsConflict('The generated total label conflicts with the selected monetary mapping.')
        if any(cell.value is not None and cell.column not in {*columns.values(), label_col} for cell in ws[row]):
            raise TotalsConflict('The generated total row contains content outside the selected monetary mappings. Review it before updating totals.')
    else:
        if any(cell.value is not None or cell.comment or cell.hyperlink for cell in ws[row]):
            raise TotalsConflict(f'GRAND TOTAL destination row {row} contains existing content. Move it to leave a blank row immediately after transactions; no rows were overwritten or shifted.')
        label_col = next((col for col in range(1, ws.max_column + 1) if col not in columns.values()), None)
        if label_col is None:
            raise TotalsConflict('There is no non-monetary column for the GRAND TOTAL label.')
    ws.cell(row, label_col, 'GRAND TOTAL')
    wb.defined_names.add(DefinedName(marker_name(ws, 'TOTAL'), attr_text=quote_sheetname(ws.title) + f'!${get_column_letter(label_col)}${row}', hidden=True))
    for key, col in columns.items():
        cell = ws.cell(row, col)
        if not generated:
            cell._style = copy(ws.cell(last, col)._style)
        cell.value = Decimal(totals['columns'][key]['total'])
        cell.number_format = CURRENCY
    for cell in ws[row]:
        font = copy(cell.font)
        font.bold = True
        cell.font = font
        border = copy(cell.border)
        border.top = Side(style='thin', color='216747')
        cell.border = border
        cell.fill = PatternFill('solid', fgColor='E6F1EC')
    totals['output'] = {'worksheet': ws.title, 'mode': 'grand_total_row', 'row': row}
    return True
