"""Workbook detection and conservative in-place-structure Excel processing."""
from copy import copy
from decimal import Decimal
from pathlib import Path
from io import BytesIO
import stat
from xml.etree import ElementTree as ET
from zipfile import ZipFile
import re
import warnings

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from gst_calculator import calculate_rate
from monetary_totals import aggregate, write_totals, header_at, TotalsConflict
from workbook_preservation import validate_export, preserve_numeric_cells


class PreservationError(ValueError):
    """Raised when OpenPyXL could alter unsupported workbook features."""


def inspect_source(source):
    """Inspect content separately from filesystem attributes; never request write access."""
    path = Path(source) if isinstance(source, (str, Path)) else None
    if path is not None and path.suffix.lower() != ".xlsx":
        raise PreservationError("Only .xlsx workbooks are supported.")
    data = path.read_bytes() if path is not None else bytes(source)
    if data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        raise PreservationError("Encrypted/password-protected Excel files and legacy binary workbooks are not supported. Upload an unencrypted .xlsx file you are authorized to use.")
    try:
        with ZipFile(BytesIO(data)) as archive:
            names = archive.namelist()
            if any("encrypted" in n.lower() or "drm" in n.lower() for n in names):
                raise PreservationError("Encrypted or rights-managed workbooks cannot be processed.")
            root = ET.fromstring(archive.read("xl/workbook.xml"))
            sharing = root.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}fileSharing")
            if sharing is not None and any(k != "readOnlyRecommended" and k != "userName" for k in sharing.attrib):
                raise PreservationError("This workbook has a password-to-modify or file-sharing security restriction. It cannot be processed without bypassing that restriction.")
    except PreservationError:
        raise
    except Exception as exc:
        raise PreservationError("Invalid or unsupported XLSX package. Upload a readable, unencrypted .xlsx workbook.") from exc
    attributes = path.stat() if path is not None else None
    windows_read_only = (bool(attributes.st_file_attributes & stat.FILE_ATTRIBUTE_READONLY)
                         if attributes is not None and hasattr(attributes, "st_file_attributes") else None)
    recommended = sharing is not None and sharing.get("readOnlyRecommended") in {"1", "true"}
    return data, {"windows_read_only": windows_read_only,
                  "read_only_recommended": recommended,
                  "read_only_detected": windows_read_only is True or recommended}


def protection_metadata(wb):
    return {"workbook_structure": bool(wb.security and wb.security.lockStructure),
            "workbook_windows": bool(wb.security and wb.security.lockWindows),
            "protected_worksheets": [ws.title for ws in wb.worksheets if ws.protection.sheet]}


def _restore_file_sharing(data, destination):
    # OpenPyXL does not model fileSharing; retain the original advisory element.
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(BytesIO(data)) as original:
        sharing = ET.fromstring(original.read("xl/workbook.xml")).find(ns + "fileSharing")
        core = original.read("docProps/core.xml") if "docProps/core.xml" in original.namelist() else None
    with ZipFile(destination) as saved:
        entries = [(entry, saved.read(entry.filename)) for entry in saved.infolist()]
    with ZipFile(destination, "w") as saved:
        for entry, content in entries:
            if entry.filename == "docProps/core.xml" and core is not None:
                content = core
            if entry.filename == "xl/workbook.xml" and sharing is not None:
                root = ET.fromstring(content)
                root.insert(1 if root.find(ns + "fileVersion") is not None else 0, sharing)
                content = ET.tostring(root, encoding="utf-8")
            saved.writestr(entry, content)


def _norm(value):
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def detect_columns(ws, scan_rows=None):
    """Find likely header cells anywhere near the top, including multirow headers."""
    candidates = {"invoice": [], "taxable": []}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, scan_rows) if scan_rows is not None else ws.max_row):
        for cell in row:
            text = _norm(cell.value)
            if header_at(ws, cell) == "invoice":
                candidates["invoice"].append((cell.row, cell.column, text))
            if header_at(ws, cell) == "taxable":
                candidates["taxable"].append((cell.row, cell.column, text))
    def choose(items, kind):
        if not items or len({item[1] for item in items}) != 1:
            return None
        # Extend the first contiguous header block through vertically split labels/units.
        ordered = sorted(items, key=lambda item: (item[0], item[1]))
        selected = ordered[0]
        for candidate in ordered[1:]:
            if candidate[0] == selected[0] + 1:
                selected = candidate
            elif candidate[0] > selected[0] + 1:
                break
        return selected
    inv, tax = choose(candidates["invoice"], "invoice"), choose(candidates["taxable"], "taxable")
    if inv is None or tax is None or inv[1] == tax[1]:
        return None
    return {"invoice_col": inv[1], "invoice_header_row": inv[0],
            "taxable_col": tax[1], "taxable_header_row": tax[0]}


def detect_rate_column(ws, scan_rows=None):
    """Return the existing Rate(%) header location, if any."""
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, scan_rows) if scan_rows is not None else ws.max_row):
        for cell in row:
            if _norm(cell.value) in {"rate", "rate percent"}:
                return cell.column, cell.row
    return None


def _cell_header(ws, col, header_row):
    # Consolidate a two-row header (for example Invoice Details / Invoice Value).
    parts = []
    for row in range(max(1, header_row - 1), header_row + 1):
        value = ws.cell(row, col).value
        if value is not None and str(value).strip():
            parts.append(str(value).strip())
    return " / ".join(dict.fromkeys(parts)) or f"Column {get_column_letter(col)}"


def _assert_supported(path, wb, ws):
    with ZipFile(path) as archive:
        names = archive.namelist()
        risky = [n for n in names if any(token in n.lower() for token in
                 ("vbaproject", "externallinks/", "pivotcache/", "slicer", "customxml/", "activex/", "embeddings/", "_xmlsignatures/"))]
        if risky or any(n.startswith("xl/drawings/") for n in names):
            raise PreservationError("This workbook contains macros, external links, pivots, drawings, embedded objects, or custom XML that OpenPyXL cannot guarantee to preserve. No output was generated.")
    if getattr(wb, "vba_archive", None):
        raise PreservationError("Macro-enabled workbooks are not supported for guaranteed preservation.")
    if any(getattr(sheet, "tables", {}) for sheet in wb.worksheets):
        raise PreservationError("Workbooks with Excel tables are not supported for guaranteed preservation.")
    if any(cell.data_type == "f" for sheet in wb.worksheets for row in sheet.iter_rows() for cell in row):
        raise PreservationError("This workbook contains formulas. OpenPyXL cannot safely update all formula references for an inserted column, so no output was generated.")
    if any(sheet.sheet_view.showGridLines is not None for sheet in wb.worksheets):
        pass


def process_workbook(source, destination, sheet_name, invoice_col, taxable_col, header_row,
                     invoice_header_row=None, taxable_header_row=None,
                     existing_rate_col=None, recalculate_existing=False, monetary_columns=None,
                     transaction_columns=None, rate_overrides=None, row_splits=None):
    """Insert Rate(%) immediately after taxable column; return summary and flagged rows."""
    rate_overrides = rate_overrides or {}
    row_splits = row_splits or {}
    destination = Path(destination)
    source_path = Path(source) if isinstance(source, (str, Path)) else None
    if source_path is not None and (source_path.resolve() == destination.resolve() or
            (destination.exists() and source_path.samefile(destination))):
        raise ValueError("The output file must not overwrite the uploaded workbook.")
    if invoice_col == taxable_col:
        raise ValueError("Invoice Value and Taxable Value must be different columns.")
    source_data, source_info = inspect_source(source)
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        wb = load_workbook(BytesIO(source_data), read_only=False, data_only=False, keep_links=True)
    if caught:
        wb.close()
        raise PreservationError("OpenPyXL reported unsupported workbook content: " + "; ".join(str(w.message) for w in caught))
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError("Selected worksheet was not found.")
    ws = wb[sheet_name]
    try:
        _assert_supported(BytesIO(source_data), wb, ws)
    except Exception:
        wb.close()
        raise
    invoice_header_row = invoice_header_row or header_row
    taxable_header_row = taxable_header_row or header_row
    detected_rate = detect_rate_column(ws)
    if existing_rate_col and (not detected_rate or existing_rate_col != detected_rate[0]):
        wb.close()
        raise PreservationError('The selected existing rate column does not have a detected Rate (%) header; no original column was overwritten.')
    existing_rate_col = existing_rate_col or (detected_rate[0] if detected_rate else None)
    if ws.protection.sheet and (not existing_rate_col or recalculate_existing or rate_overrides):
        wb.close()
        raise PreservationError("The selected worksheet is protected. Updating Rate(%) would bypass worksheet protection. Use an authorized unprotected copy or keep existing rates.")
    totals = aggregate(ws, header_row + 1, invoice_col, taxable_col, monetary_columns, transaction_columns)
    if detected_rate and any(item['source_column'] == detected_rate[0] for item in totals['columns'].values()):
        wb.close()
        raise ValueError('Rate (%) cannot be a monetary column.')
    included_rows = set(totals['included_rows'])
    rate_col = existing_rate_col or (taxable_col + 1)

    rate_letter = get_column_letter(rate_col)
    if not existing_rate_col:
        # OpenPyXL does not shift merged ranges when inserting columns. Rebuild them
        # explicitly so the surrounding title/layout remains intact.
        if ws.data_validations.count or ws.conditional_formatting:
            wb.close()
            raise PreservationError('Data validation or conditional formatting references cannot be safely shifted for an inserted rate column. Use a workbook with an existing Rate (%) column or an authorized simplified copy.')
        for name in wb.defined_names.values():
            if not name.name.startswith('_GST_LEDGER_'):
                wb.close()
                raise PreservationError('Named references cannot be safely shifted for an inserted rate column. Use an existing Rate (%) column or review the workbook.')
        merged = list(ws.merged_cells.ranges)
        saved_styles = {(cell.row, cell.column): copy(cell._style) for cell in ws._cells.values() if cell.has_style}
        freeze = ws.freeze_panes
        print_area = list(ws._print_area.ranges) if ws.print_area else []
        print_columns = ws.print_title_cols
        column_dimensions = [copy(dimension) for dimension in ws.column_dimensions.values()]
        taxable_width = next((dimension.width for dimension in column_dimensions
                              if dimension.min <= taxable_col <= dimension.max),
                             ws.sheet_format.defaultColWidth or 13)
        filter_ref = ws.auto_filter.ref
        for merged_range in merged:
            ws.unmerge_cells(str(merged_range))
        ws.insert_cols(rate_col, 1)
        # insert_cols leaves column dimension keys and filter ranges at their old
        # addresses. Shift them explicitly so widths and filters stay with columns.
        ws.column_dimensions.clear()
        for dimension in column_dimensions:
            first, last = dimension.min, dimension.max
            if first >= rate_col:
                first += 1
                last += 1
            elif last >= rate_col:
                last += 1
            dimension.min, dimension.max = first, last
            dimension.index = get_column_letter(first)
            ws.column_dimensions[dimension.index] = dimension
        if filter_ref:
            shifted_ranges = []
            for area in filter_ref.split():
                min_col, min_row, max_col, max_row = range_boundaries(area)
                if min_col >= rate_col:
                    min_col += 1
                    max_col += 1
                elif max_col >= rate_col - 1:
                    max_col += 1
                shifted_ranges.append(f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}")
            ws.auto_filter.ref = " ".join(shifted_ranges)
        for merged_range in merged:
            min_col, min_row, max_col, max_row = merged_range.bounds
            if min_col >= rate_col:
                min_col += 1
                max_col += 1
            elif max_col >= rate_col:
                max_col += 1
            ws.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)
        # Re-merging can redistribute borders; restore every original styled cell.
        for (row, col), style in saved_styles.items():
            ws.cell(row, col + (col >= rate_col))._style = style
        if freeze:
            from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
            letter, number = coordinate_from_string(freeze)
            column = column_index_from_string(letter)
            shifted_column = column + (column >= rate_col)
            ws.sheet_view.pane.topLeftCell = f'{get_column_letter(shifted_column)}{number}'
            if shifted_column != column:
                ws.sheet_view.pane.xSplit = shifted_column - 1
        if print_area:
            shifted = []
            for area in print_area:
                first, top, last, bottom = area.bounds
                shifted.append(f'{get_column_letter(first + (first >= rate_col))}{top}:{get_column_letter(last + (last >= rate_col))}{bottom}')
            ws.print_area = shifted
        if print_columns:
            from openpyxl.utils.cell import column_index_from_string
            first, last = [column_index_from_string(part.replace('$', '')) for part in print_columns.split(':')]
            ws.print_title_cols = f'{get_column_letter(first + (first >= rate_col))}:{get_column_letter(last + (last >= rate_col))}'
        for cell in ws._cells.values():
            if cell.hyperlink:
                cell.hyperlink.ref = cell.coordinate
        from openpyxl.cell.cell import MergedCell
        if isinstance(ws.cell(taxable_header_row, rate_col), MergedCell) or any(
                isinstance(ws.cell(row, rate_col), MergedCell) for row in included_rows):
            wb.close()
            raise PreservationError('The rate column intersects a merged header or transaction cell. No source data was overwritten.')
        ws.cell(taxable_header_row, rate_col).value = "Rate(%)"
        source_style = ws.cell(taxable_header_row, taxable_col)
        target = ws.cell(taxable_header_row, rate_col)
        if source_style.has_style:
            target._style = copy(source_style._style)
        if source_style.number_format:
            target.number_format = source_style.number_format
        # Match the new field's width to the taxable value column.
        # Looking up an absent dimension key creates a default-width override,
        # which can hide an existing grouped column's width. Use the source range.
        ws.column_dimensions[rate_letter].width = taxable_width
    if not existing_rate_col and invoice_col >= rate_col:
        invoice_col += 1
    try:
        totals_changed = write_totals(wb, ws, totals, rate_col if not existing_rate_col else None)
    except TotalsConflict as exc:
        wb.close()
        raise PreservationError(str(exc)) from exc
    summary = {0: 0, 5: 0, 12: 0, 18: 0, 28: 0, "unmatched": 0}
    flagged = []
    inv_header = _cell_header(ws, invoice_col, invoice_header_row)
    tax_header = _cell_header(ws, taxable_col, taxable_header_row)
    # Keep original physical row numbering and blank rows; process each row independently.
    for row in sorted(included_rows):
        result = calculate_rate(ws.cell(row, invoice_col).value, ws.cell(row, taxable_col).value)
        if not existing_rate_col or recalculate_existing or row in rate_overrides:
            ws.cell(row, rate_col).value = rate_overrides.get(row, result.rate)
        if result.rate is None:
            summary["unmatched"] += 1
        else:
            summary[result.rate] += 1
        if result.flagged:
            flagged.append({"row": row, "invoice_letter": get_column_letter(invoice_col), "invoice_col": invoice_col,
                            "invoice_header": inv_header, "taxable_letter": get_column_letter(taxable_col),
                            "taxable_col": taxable_col, "taxable_header": tax_header,
                            "rate_letter": rate_letter, "rate_col": rate_col,
                            "difference": str(result.difference) if result.difference is not None else "",
                            "rate": result.rate, "message": result.message})
    try:
        if row_splits:
            from row_splits import transform_rows
            try:
                transform_rows(wb, ws, row_splits, totals, rate_col, None if existing_rate_col else rate_col,
                               transaction_columns or {})
            except ValueError as exc:
                raise PreservationError(str(exc)) from exc
        if existing_rate_col and not recalculate_existing and not totals_changed and not rate_overrides and not row_splits:
            # Preserve the original bytes when no workbook update is possible or needed.
            destination.write_bytes(source_data)
        else:
            wb.save(destination)
            _restore_file_sharing(source_data, destination)
        # Preserve source numeric precision before reopening, then verify every
        # original cell normally. Numeric serialization loss is not a tolerated mismatch.
        baseline = load_workbook(BytesIO(source_data), read_only=False, data_only=False)
        check = None
        try:
            preserve_numeric_cells(source_data, destination, baseline, sheet_name,
                                   None if existing_rate_col else rate_col, rate_col,
                                   recalculate_existing, included_rows, totals, rate_overrides, row_splits)
            check = load_workbook(destination, read_only=False, data_only=False)
            if check.sheetnames != wb.sheetnames or rate_col > check[sheet_name].max_column:
                raise PreservationError('Export validation failed; output was removed.')
            validate_export(baseline, check, wb, sheet_name, None if existing_rate_col else rate_col,
                            rate_col, recalculate_existing, included_rows, totals, rate_overrides, row_splits)
        except ValueError as exc:
            raise PreservationError(str(exc)) from exc
        finally:
            baseline.close()
            if check is not None:
                check.close()
    except Exception:
        if destination.exists():
            destination.unlink()
        raise
    finally:
        wb.close()
    return {"source_info": source_info, "summary": summary, "flagged": flagged, "rate_col": rate_col,
            "rate_letter": rate_letter, "invoice_header": inv_header, "taxable_header": tax_header,
            "column_totals": totals}
